# -*- coding: utf-8 -*-
"""Generează documentul de documentare a SA3.1 (informare și publicitate) pentru RP1,
după modelul „SA_6.1_06.2026_Informare și Publicitate” din proiectul PIDS, pe șablonul oficial
„Antet_subsol_UNIC_Porolissum_PEO_A4_portret.docx” (antet + subsol, de folosit la toate documentele proiectului).
Textele între «…» sunt completări de făcut de echipa de proiect (evidențiate cu galben)."""
import os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
LIV = os.path.join(HERE, 'livrabile')
IMG = os.path.join(LIV, 'sa31_img')
DATA_CAPTURI = '07.10.2026'

SABLON = os.path.join(LIV, 'Antet_subsol_UNIC_Porolissum_PEO_A4_portret.docx')  # antetul oficial (din 07.10.2026)
doc = None; CONTENT_W = None

def document_nou():
    global doc, CONTENT_W
    doc = Document(SABLON)
    body = doc.element.body
    for el in list(body):
        if not el.tag.endswith('}sectPr'):
            body.remove(el)
    sec = doc.sections[0]  # marginile, antetul și subsolul rămân cele din șablon
    CONTENT_W = sec.page_width - sec.left_margin - sec.right_margin
    st = next(x for x in doc.styles if x.type == 1 and x.name.lower() == 'normal')
    st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

def _runs(p, text, bold=False, size=None, color=None, italic=False):
    """scrie textul, evidențiind cu galben fragmentele «…»"""
    for part in re.split(r'(«[^»]*»)', text):
        if not part:
            continue
        r = p.add_run(part)
        r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor(*color)
        if part.startswith('«'):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p

def P(text='', bold=False, size=None, align=None, space_after=6, italic=False, color=None):
    p = doc.add_paragraph()
    if align == 'center': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'justify': p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    _runs(p, text, bold=bold, size=size, italic=italic, color=color)
    return p

def LINK(url, text=None, bullet=False, bold=True):
    p = doc.add_paragraph()
    if bullet:  # șablonul nu are stilul „List Bullet”: marcator manual
        p.paragraph_format.left_indent = Cm(0.63); p.paragraph_format.first_line_indent = Cm(-0.63)
        p.add_run('•\t')
    part = p.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), r_id)
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '0563C1'); rPr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    if bold:
        rPr.append(OxmlElement('w:b'))
    r.append(rPr)
    t = OxmlElement('w:t'); t.text = text or url; t.set(qn('xml:space'), 'preserve'); r.append(t)
    h.append(r); p._p.append(h)
    p.paragraph_format.space_after = Pt(6)
    return p

def TITLU(text, size=20, before=180):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(12)
    _runs(p, text, bold=True, size=size)
    return p

def PAGEBREAK():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def IMAGINE(fname, width=None, legenda=None, max_h_cm=19.5):
    from PIL import Image
    path = os.path.join(IMG, fname)
    w, h = Image.open(path).size
    width = width or CONTENT_W
    if h / w * width > Cm(max_h_cm):
        width = int(Cm(max_h_cm) * w / h)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=width)
    p.paragraph_format.space_after = Pt(2)
    if legenda:
        P(legenda, size=10, italic=True, align='center', space_after=10, color=(0x40, 0x40, 0x40))

def PLACEHOLDER(text):
    """casetă pentru o captură/fotografie care trebuie adăugată"""
    t = doc.add_table(rows=1, cols=1)
    cell = t.rows[0].cells[0]
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + side); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '6'); e.set(qn('w:color'), '808080'); b.append(e)
    tcPr.append(b)
    cell.width = CONTENT_W
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60); p.paragraph_format.space_after = Pt(60)
    _runs(p, '«' + text + '»', size=11, italic=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


LUNI = {'07': ('iulie 2026', 'L1'), '08': ('august 2026', 'L2'), '09': ('septembrie 2026', 'L3')}
FB_URL = 'https://www.facebook.com/unic.viitorprineducatie?locale=ro_RO'
IG_URL = 'https://www.instagram.com/«cont_instagram»/'
SITE = 'https://napocaporolissum.ro/unic-porolissum-2/'

def titlu_pagina(luna, cuprins):
    nume, L = LUNI[luna]
    TITLU('S.A.3.1 – Derularea activităților de informare și publicitate a proiectului', size=20, before=150)
    P('Luna de raportare: ' + nume + ' (' + L + ')', bold=True, size=14, align='center', space_after=4)
    P('Raportul de progres nr. 1 (iulie – septembrie 2026)', size=13, align='center', space_after=40)
    P('Proiect: UNIC – „Uniți pentru Nevoile Incluzive și Continuitatea Educației Elevilor”, cod SMIS 352704', size=11, align='center', space_after=2)
    P('Beneficiar: Asociația Grupul de Acțiune Locală Napoca Porolissum', size=11, align='center', space_after=2)
    P('Contract de finanțare nr. 11205/16.06.2026 – Programul Educație și Ocupare 2021-2027', size=11, align='center', space_after=60)
    P('Conținut:', bold=True, size=11, space_after=2)
    for i, item in enumerate(cuprins, 1):
        P(f'{i}. {item}', size=11, space_after=1)
    P('', space_after=40)
    P('Întocmit: Asistent manager, «nume prenume»', size=11, space_after=2)
    P('Avizat: Manager de proiect, Baba Alina-Ioana', size=11, space_after=2)
    P('Capturile de ecran ale site-ului au fost realizate la ' + DATA_CAPTURI + ', cu adresa URL activă; capturile postărilor se realizează din conturile proiectului, cu data vizibilă.', size=10, italic=True, space_after=2)

def sectiune_postari(retea, url, luna, nr=4, ultima=None):
    """pagini rezervate pentru postările unei rețele sociale într-o lună"""
    nume = LUNI[luna][0]
    PAGEBREAK()
    TITLU('POSTĂRI PE ' + ('PAGINA DE FACEBOOK' if retea == 'fb' else 'CONTUL DE INSTAGRAM') + ' A' + ('' if retea == 'fb' else 'L') + ' PROIECTULUI', size=16, before=120)
    P('„UNIC: Viitor prin Educație”' if retea == 'fb' else '«@cont_instagram»', bold=True, size=14, align='center', space_after=4)
    LINK(url)
    P('Postări publicate în luna ' + nume + ': «__». Pentru fiecare postare: link-ul (Distribuie → Copiază link) și captura de ecran, cu data vizibilă, în ordine cronologică.', align='justify', size=11)
    for i in range(1, nr + 1):
        PAGEBREAK()
        if ultima and i == nr:  # ultima postare cunoscută din lună (ordine cronologică)
            data, tema, link = ultima
            LINK(link, '«link postare: https://www.facebook.com/share/p/…»', bullet=True)
            P(data + ' – ' + tema, size=11)
        else:
            LINK(url, f'«Postarea {i} – link»', bullet=True)
            P(f'«data postării» – «tema postării {i}»', size=11)
        PLACEHOLDER(f'captură de ecran postarea {i} – {nume}')

def sectiune_afis(unde, foto):
    PAGEBREAK()
    TITLU('EXPUNEREA AFIȘULUI A3 AL PROIECTULUI', size=16, before=120)
    P(unde, align='justify', size=11)
    for f in foto:
        PLACEHOLDER(f)

def build_luna(luna):
    document_nou()
    nume = LUNI[luna][0]
    out = os.path.join(LIV, f'SA 3.1_{luna}.2026_Informare și Publicitate.docx')
    if luna == '07':
        titlu_pagina(luna, ['Subpagina proiectului pe site-ul beneficiarului (creată la 02.07.2026)',
                            'Anunțul de începere a proiectului',
                            'Metodologiile publicate pe site (16.07.2026)',
                            'Materiale de informare realizate (afiș A3, flyere)',
                            'Expunerea afișului A3 la sediul beneficiarului',
                            'Postări pe pagina de Facebook a proiectului',
                            'Postări pe contul de Instagram al proiectului'])
        PAGEBREAK()
        TITLU('SUBPAGINA PROIECTULUI PE SITE-UL BENEFICIARULUI', size=16, before=120)
        LINK(SITE)
        P('Subpagina dedicată proiectului a fost creată la 02.07.2026, în secțiunea „Programe – PEO – Programul Educație și Ocupare”, cu antetul proiectului '
          '(emblema Uniunii Europene și mențiunea „Cofinanțat de Uniunea Europeană”, sigla Guvernului României, sigla beneficiarului și sigla proiectului). '
          'Prezintă descrierea și obiectivul principal, justificarea, rezultatele așteptate, activitățile, anunțurile și metodologiile publicate.', align='justify', size=11)
        PAGEBREAK(); LINK(SITE, bullet=True)
        IMAGINE('site_header.png', legenda='Partea superioară a subpaginii – coperta și antetul proiectului (captură ' + DATA_CAPTURI + ')')
        IMAGINE('site_acc_despre.png', legenda='Secțiunea „Despre proiectul UNIC”')
        PAGEBREAK(); LINK(SITE, bullet=True)
        IMAGINE('site_acc_rezultate.png', legenda='Secțiunea „Rezultatele proiectului UNIC”')
        PAGEBREAK()
        TITLU('ANUNȚUL DE ÎNCEPERE A PROIECTULUI', size=16, before=120)
        LINK(SITE)
        P('Anunțul de începere a fost publicat pe subpagina proiectului la data de «__.07.2026» și cuprinde elementele prevăzute de Manualul de identitate vizuală 2021-2027: '
          'beneficiarul, titlul și codul MySMIS al proiectului, programul și prioritatea, fondul (FSE+), perioada și regiunea de implementare, valoarea totală, '
          'scopul, activitățile, grupul țintă, rezultatele urmărite și datele de contact.', align='justify', size=11)
        PAGEBREAK(); LINK(SITE, bullet=True)
        IMAGINE('site_acc_anunturi.png', legenda='Anunțul de începere a proiectului – secțiunea „Anunțuri” (captură ' + DATA_CAPTURI + ')')
        P('«Se atașează și PDF-ul anunțului de începere, în forma avizată de Managerul de proiect.»', size=10, italic=True)
        PAGEBREAK()
        TITLU('METODOLOGIILE PUBLICATE PE SITE', size=16, before=120)
        LINK(SITE)
        P('La 16.07.2026 au fost publicate pe subpagina proiectului metodologiile de selecție a grupului țintă și de acordare a sprijinului financiar, '
          'asigurând accesul egal la informație al tuturor persoanelor interesate.', align='justify', size=11)
        IMAGINE('site_acc_metodologii.png', legenda='Secțiunea „Metodologii” (captură ' + DATA_CAPTURI + ')')
        LINK('https://napocaporolissum.ro/wp-content/uploads/METODOLOGIE_GT_CU_ANEXE_UNIC_352704_16.07.2026.pdf',
             'Metodologia privind identificarea, recrutarea, înscrierea, verificarea, selectarea, validarea și menținerea grupului țintă (16.07.2026)', bullet=True, bold=False)
        LINK('https://napocaporolissum.ro/wp-content/uploads/METODOLOGIE_SPRIJIN_FINANCIAR_CU_ANEXE_UNIC_352704_16.07.2026.pdf',
             'Metodologia privind acordarea sprijinului financiar pentru transport, cazare și masă (16.07.2026)', bullet=True, bold=False)
        PAGEBREAK()
        TITLU('MATERIALE DE INFORMARE REALIZATE', size=16, before=120)
        P('Materialele au fost realizate în luna iulie 2026 (versiuni finale din 31.07.2026), cu elementele obligatorii de identitate vizuală, '
          'și au fost utilizate în întâlnirile de informare cu unitățile de învățământ partenere și cu elevii.', align='justify', size=11)
        P('• Afiș A3 al proiectului (fișier: „Afiș A3_proiect UNIC2.pdf”);', size=11, space_after=1)
        P('• Flyer de prezentare a proiectului pentru unitățile de învățământ, față-verso (fișier: „flyer_UNIC_352704_scoli_fata_verso.pdf”);', size=11, space_after=1)
        P('• Flyer de prezentare a proiectului, format 12 × 17 cm, față-verso (fișier: „Flyer-UNIC2 (12 x 17 cm).pdf”).', size=11, space_after=8)
        PAGEBREAK(); P('Afiș A3 al proiectului', bold=True, size=12, align='center')
        PLACEHOLDER('imagine afiș A3 – se inserează din fișierul „Afiș A3_proiect UNIC2.pdf”')
        PAGEBREAK(); P('Flyer de prezentare pentru unitățile de învățământ (față)', bold=True, size=12, align='center'); IMAGINE('flyer_scoli-1.png', max_h_cm=21)
        PAGEBREAK(); P('Flyer de prezentare pentru unitățile de învățământ (verso)', bold=True, size=12, align='center'); IMAGINE('flyer_scoli-2.png', max_h_cm=21)
        PAGEBREAK(); P('Flyer 12 × 17 cm (față și verso)', bold=True, size=12, align='center')
        t = doc.add_table(rows=1, cols=2)
        for cell, f in zip(t.rows[0].cells, ['flyer_12x17-1.png', 'flyer_12x17-2.png']):
            cp = cell.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.add_run().add_picture(os.path.join(IMG, f), width=Cm(7.6))
        sectiune_afis('Afișul A3 al proiectului a fost expus, la data de «__.07.2026», într-un loc vizibil publicului la sediul de implementare al beneficiarului din comuna Gilău, '
                      'conform cerințelor Ghidului de identitate vizuală (Anexa 18 la Manualul beneficiarului).',
                      ['fotografie – afișul A3 expus la sediul beneficiarului (Gilău)'])
        sectiune_postari('fb', FB_URL, luna, nr=4)
        sectiune_postari('ig', IG_URL, luna, nr=3)
    elif luna == '08':
        titlu_pagina(luna, ['Expunerea afișului A3 la unitățile de învățământ partenere',
                            'Postări pe pagina de Facebook a proiectului',
                            'Postări pe contul de Instagram al proiectului'])
        sectiune_afis('Afișul A3 al proiectului a fost expus, la data de «__.08.2026», într-un loc vizibil publicului la unitățile de învățământ partenere: '
                      'Colegiul Tehnic Turda și Liceul Tehnologic „Vlădeasa” Huedin.',
                      ['fotografie – afișul A3 expus la Colegiul Tehnic Turda', 'fotografie – afișul A3 expus la Liceul Tehnologic „Vlădeasa” Huedin'])
        sectiune_postari('fb', FB_URL, luna, nr=4)
        sectiune_postari('ig', IG_URL, luna, nr=3)
    else:
        titlu_pagina(luna, ['Pagina de Facebook a proiectului – situația la finalul perioadei de raportare',
                            'Postări pe pagina de Facebook a proiectului',
                            'Postări pe contul de Instagram al proiectului',
                            'Expunerea afișului A3 la Liceul Teologic Reformat Cluj-Napoca',
                            'Inscripționarea echipamentelor IT'])
        PAGEBREAK()
        TITLU('PAGINA DE FACEBOOK A PROIECTULUI', size=16, before=120)
        P('„UNIC: Viitor prin Educație”', bold=True, size=14, align='center', space_after=4)
        LINK(FB_URL)
        P('Pagina proiectului a fost creată la «__.__.2026» și avea, la data capturii (' + DATA_CAPTURI + '), 588 de urmăritori. '
          'Secțiunea „Prezentare” conține titlul proiectului, mențiunea cofinanțării de către Uniunea Europeană prin Programul Educație și Ocupare (PEO 2021-2027), '
          'codul MySMIS 352704, adresa de e-mail și trimiterea către subpagina proiectului de pe site-ul beneficiarului.', align='justify', size=11)
        IMAGINE('fb_pagina.png', legenda='Pagina de Facebook a proiectului – antet, secțiunea „Prezentare” și postarea din 29.09.2026 (captură ' + DATA_CAPTURI + ')')
        sectiune_postari('fb', FB_URL, luna, nr=5,
                         ultima=('29.09.2026', '„Au început activitățile proiectului UNIC la Colegiul Tehnic Turda!” – postare cu fotografii și video din cadrul activităților cu elevii', FB_URL))
        sectiune_postari('ig', IG_URL, luna, nr=3)
        sectiune_afis('Afișul A3 al proiectului a fost expus, la data de «__.09.2026», la Liceul Teologic Reformat Cluj-Napoca, unitate de învățământ parteneră '
                      'începând cu acordul de colaborare din 09.09.2026.',
                      ['fotografie – afișul A3 expus la Liceul Teologic Reformat Cluj-Napoca'])
        PAGEBREAK()
        TITLU('INSCRIPȚIONAREA ECHIPAMENTELOR IT', size=16, before=120)
        P('Echipamentele IT recepționate la 18.09.2026 (3 laptopuri HP Envy 17, 1 laptop Acer Nitro V15, 2 desktopuri și 2 multifuncționale Epson) '
          '«au fost inscripționate la data de __.__.2026» cu autocolante conținând elementele de identitate vizuală ale proiectului.', align='justify', size=11)
        PLACEHOLDER('fotografii – autocolantele aplicate pe echipamentele IT')
    doc.save(out)
    print('OK ->', out)

for _luna in ('07', '08', '09'):
    build_luna(_luna)
