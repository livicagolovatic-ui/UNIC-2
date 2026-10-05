# -*- coding: utf-8 -*-
"""Motor de generare pentru Raportul de progres nr. 1 – UNIC (SMIS 352704).

Fiecare câmp MySMIS se scrie cu FIELD(eticheta, limita, text). Textul este
cel care se copiază în MySMIS; caracterele se numără exact (cu spații și
rânduri noi) și se compară cu limita afișată de contorul MySMIS.
Marcajul «...» = informație de confirmat/completat (evidențiată cu galben).
"""
import re, sys
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = docx.Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.left_margin = s.right_margin = Cm(2.0)
s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)

INK = RGBColor(0x1A, 0x1A, 0x1A)
BLUE = RGBColor(0x00, 0x3D, 0x82)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
RED = RGBColor(0x9B, 0x1C, 0x1C)
GREEN = RGBColor(0x1E, 0x6B, 0x2E)

st = doc.styles
n = st['Normal']
n.font.name = 'Calibri'; n.font.size = Pt(10.5); n.font.color.rgb = INK
n._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
n.paragraph_format.space_after = Pt(4)
n.paragraph_format.line_spacing = 1.1


def _cfg(name, size, bold, color, before, after):
    x = st[name]
    x.font.name = 'Calibri'; x.font.size = Pt(size); x.font.bold = bold
    x.font.color.rgb = color
    x.paragraph_format.space_before = Pt(before)
    x.paragraph_format.space_after = Pt(after)
    x.paragraph_format.keep_with_next = True


_cfg('Heading 1', 16, True, BLUE, 18, 8)
_cfg('Heading 2', 13, True, BLUE, 14, 6)
_cfg('Heading 3', 11.5, True, INK, 10, 4)
_cfg('Heading 4', 10.5, True, GREY, 8, 3)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)


def borders(tbl, color='BFBFBF', sz=4):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        b.append(e)
    tblPr.append(b)


TOK = re.compile(r'(\*\*.+?\*\*|«.+?»)', re.S)


def runs(p, text, size=None, color=None, italic=False, bold=False):
    for part in TOK.split(text):
        if not part:
            continue
        b = bold; hl = False
        if part.startswith('**') and part.endswith('**'):
            part, b = part[2:-2], True
        elif part.startswith('«') and part.endswith('»'):
            hl = True
        r = p.add_run(part); r.bold = b; r.italic = italic
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color
        if hl: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def H1(t): doc.add_paragraph(t, style='Heading 1')
def H2(t): doc.add_paragraph(t, style='Heading 2')
def H3(t): doc.add_paragraph(t, style='Heading 3')
def H4(t): doc.add_paragraph(t, style='Heading 4')


def P(t, size=None, color=None, italic=False):
    p = doc.add_paragraph(); runs(p, t, size=size, color=color, italic=italic)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def SMALL(t):
    p = doc.add_paragraph(); runs(p, t, size=8.5, color=GREY, italic=True)
    return p


def BUL(items, size=None):
    for it in items:
        p = doc.add_paragraph(style='List Bullet'); runs(p, it, size=size)
        p.paragraph_format.space_after = Pt(2)


def FLAG(t):
    """Notă internă a managerului de proiect – NU se copiază în MySMIS."""
    tbl = doc.add_table(rows=1, cols=1); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0); shade(c, 'FBEAEA'); borders(tbl, 'E0A0A0', 4)
    p = c.paragraphs[0]
    runs(p, '⚠ Notă MP (nu se copiază): ', size=9, color=RED, bold=True)
    runs(p, t, size=9, color=RED)
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def CALLOUT(title, body, fill='FFF6E5', line='E8C97A'):
    tbl = doc.add_table(rows=1, cols=1); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0); shade(c, fill); borders(tbl, line, 6)
    p = c.paragraphs[0]; runs(p, title, size=10, bold=True)
    for para in body.split('\n'):
        q = c.add_paragraph(); runs(q, para, size=9.5)
        q.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def TBL(rows, widths=None, small=True, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci); c.text = ''
            first = True
            for line in str(val).split('\n'):
                tgt = c.paragraphs[0] if first else c.add_paragraph()
                tgt.paragraph_format.space_before = Pt(1); tgt.paragraph_format.space_after = Pt(1)
                runs(tgt, line, size=8.5 if small else 9.5, bold=(header and ri == 0),
                     color=BLUE if (header and ri == 0) else None)
                first = False
            if header and ri == 0:
                shade(c, 'E8EEF7')
    if widths:
        for ri in range(len(rows)):
            for ci, w in enumerate(widths):
                t.cell(ri, ci).width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def PAGEBREAK():
    doc.add_page_break()


# ---------------------------------------------------------------- câmpuri
FIELDS = []          # (secțiune, etichetă, limită, n caractere, text)
SECTION = {'t': ''}


def SECT(t):
    SECTION['t'] = t
    H2(t)


def plain(text):
    """Textul exact care se lipește în MySMIS (fără marcaje de formatare)."""
    return text.replace('**', '').strip()


def FIELD(label, limit, text, note=None):
    txt = plain(text)
    n = len(txt)
    FIELDS.append((SECTION['t'], label, limit, n, txt))
    # eticheta câmpului
    tbl = doc.add_table(rows=1, cols=1); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0); shade(c, 'E8EEF7'); borders(tbl, 'C7D6EA', 4)
    p = c.paragraphs[0]; runs(p, label, size=9.5, color=BLUE, bold=True)
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
    # textul
    for para in text.strip().split('\n'):
        if not para.strip():
            continue
        if para.startswith('• '):
            q = doc.add_paragraph(style='List Bullet'); runs(q, para[2:])
            q.paragraph_format.space_after = Pt(1)
        else:
            q = doc.add_paragraph(); runs(q, para)
            q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # contorul
    pct = 100.0 * n / limit
    ok = n <= limit * 0.95
    q = doc.add_paragraph()
    runs(q, 'Caractere: %s / %s (%.0f%%) – %s' % (f'{n:,}'.replace(',', '.'),
         f'{limit:,}'.replace(',', '.'), pct,
         'în limită' if ok else ('ATENȚIE: aproape de limită' if n <= limit else 'DEPĂȘEȘTE LIMITA')),
         size=8, color=GREEN if ok else RED, italic=True)
    q.paragraph_format.space_after = Pt(8)
    if note:
        FLAG(note)


def report():
    print('%-6s %-7s %s' % ('car.', 'limită', 'câmp'), file=sys.stderr)
    for sec, lab, lim, n, _ in FIELDS:
        flag = '!!' if n > lim * 0.95 else '  '
        print('%s %6d / %5d  %s | %s' % (flag, n, lim, sec[:28], lab[:70]), file=sys.stderr)


def export_txt(path):
    with open(path, 'w', encoding='utf-8') as fh:
        fh.write('RAPORT DE PROGRES NR. 1 – UNIC (SMIS 352704)\n')
        fh.write('Text pentru copiere în MySMIS2021. Marcajul «...» = de confirmat/completat înainte de copiere.\n')
        fh.write('=' * 78 + '\n')
        last = None
        for sec, lab, lim, n, txt in FIELDS:
            if sec != last:
                fh.write('\n\n' + '#' * 78 + '\n' + sec + '\n' + '#' * 78 + '\n')
                last = sec
            fh.write('\n--- %s  [%d / %d caractere]\n\n' % (lab, n, lim))
            fh.write(txt + '\n')
