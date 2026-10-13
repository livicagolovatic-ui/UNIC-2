# -*- coding: utf-8 -*-
"""Generează „livrabile/RP1_texte_SA1.1_A4_A5.docx”: textele de progres (Rezultat / Abateri / Progres) pentru
SA1.1, SA4.1, SA4.2, SA5.1, SA5.2, SA5.3, SA5.4, pe antetul oficial, cu numărul de caractere față de limita MySMIS.
Textele între «…» sunt completări de făcut (galben)."""
import os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
LIV = os.path.join(HERE, 'livrabile')
OUT = os.path.join(LIV, 'RP1_texte_SA1.1_A4_A5.docx')
FISIERE = ['04_SA1.1_management', '04_SA4.1_selectie', '04_SA4.2_identificare', '04_SA5.1_sprijin',
           '04_SA5.2_formare', '04_SA5.3_informare', '04_SA5.4_remediale']

doc = Document(os.path.join(LIV, 'Antet_subsol_UNIC_Porolissum_PEO_A4_portret.docx'))
body = doc.element.body
for el in list(body):
    if not el.tag.endswith('}sectPr'):
        body.remove(el)
st = next(x for x in doc.styles if x.type == 1 and x.name.lower() == 'normal')
st.font.name = 'Times New Roman'; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

def runs(p, text, bold=False, size=None, italic=False, color=None):
    for part in re.split(r'(«[^»]*»)', text):
        if not part: continue
        r = p.add_run(part); r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor(*color)
        if part.startswith('«'): r.font.highlight_color = WD_COLOR_INDEX.YELLOW

def P(text, bold=False, size=None, align=None, after=6, italic=False, color=None, before=0):
    p = doc.add_paragraph()
    if align == 'center': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'justify': p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    runs(p, text, bold, size, italic, color); return p

P('Raportul de progres nr. 1 (iulie – septembrie 2026) – texte pentru secțiunea 4 „Activități implementate și rezultate obținute”',
  bold=True, size=14, align='center', before=24, after=4)
P('SA1.1 (actualizată cu întâlnirea din 25.09.2026), SA4.1, SA4.2, SA5.1, SA5.2, SA5.3, SA5.4', size=11, align='center', after=4)
P('Proiect UNIC – „Uniți pentru Nevoile Incluzive și Continuitatea Educației Elevilor”, cod SMIS 352704 | versiunea din 13.10.2026', size=10, align='center', after=14)
P('Fiecare subactivitate are trei câmpuri, în ordinea din MySMIS2021: Rezultat obținut în perioada de raportare (limită 3.500 de caractere), Abateri/riscuri identificate și Progres în perioada de raportare (limită 10.500 de caractere). '
  'Numărul de caractere este indicat sub fiecare titlu. Marcajele galbene «…» sunt completări de făcut înainte de transmitere; textele se copiază în MySMIS fără formatare.',
  size=10, italic=True, align='justify', after=10)
first = True
for f in FISIERE:
    s = open(os.path.join(HERE, 'sectiuni', f + '.txt'), encoding='utf-8').read()
    if not first:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    first = False
    for m in re.finditer(r'=== (.+?)\n(.*?)(?=\n=== |\Z)', s, re.S):
        titlu, txt = m.group(1), m.group(2).strip()
        lim = 3500 if 'Rezultat' in titlu else 10500
        P(titlu, bold=True, size=12, before=10, after=2)
        P(f'{len(txt):,} / {lim:,} caractere'.replace(',', '.'), size=9, italic=True, color=(0x59, 0x59, 0x59), after=4)
        for par in txt.split('\n'):
            P(par, align='justify', after=4)
doc.save(OUT); print('OK ->', OUT)
