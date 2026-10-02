# -*- coding: utf-8 -*-
"""Motor comun pentru documentele Word ale proiectului Intervenția 5.

Aceleași funcții ca docx_engine.py din rădăcina repo-ului (H1, P, TBL, CALLOUT ...),
cu paleta GAL Napoca Porolissum (maro #3D2E31, verde #4FAB50) și subsol cu numerotare.
"""
import docx, re
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = docx.Document()

s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.left_margin = s.right_margin = Cm(2.0)
s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)

INK   = RGBColor(0x1A, 0x1A, 0x1A)
BROWN = RGBColor(0x3D, 0x2E, 0x31)   # maro GAL
GREEN = RGBColor(0x2F, 0x6B, 0x30)   # verde GAL închis, lizibil pe alb
GREY  = RGBColor(0x5A, 0x5A, 0x5A)
RED   = RGBColor(0x9B, 0x1C, 0x1C)
HEAD_FILL = 'E6F2E6'                 # verde GAL foarte deschis
HEAD_LINE = 'A9D3AA'

st = doc.styles
n = st['Normal']
n.font.name = 'Calibri'; n.font.size = Pt(10.5); n.font.color.rgb = INK
n._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
n.paragraph_format.space_after = Pt(6)
n.paragraph_format.line_spacing = 1.12

def cfg(name, size, bold, color, before, after, keep=True):
    x = st[name]
    x.font.name = 'Calibri'; x.font.size = Pt(size); x.font.bold = bold
    x.font.color.rgb = color
    x.paragraph_format.space_before = Pt(before)
    x.paragraph_format.space_after = Pt(after)
    x.paragraph_format.keep_with_next = keep
    return x

cfg('Title', 20, True, BROWN, 0, 4)
cfg('Heading 1', 16, True, BROWN, 18, 6)
cfg('Heading 2', 13, True, GREEN, 14, 5)
cfg('Heading 3', 11.5, True, INK, 10, 3)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)

def borders(tbl, color='BFBFBF', sz=4):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        b.append(e)
    tblPr.append(b)

TBLPR_ORDER = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize',
               'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook']

def order_tblPr(t):
    """Word cere copiii lui w:tblPr în ordinea din schemă."""
    tblPr = t._tbl.tblPr
    kids = list(tblPr)
    key = lambda e: TBLPR_ORDER.index(e.tag.split('}')[1]) if e.tag.split('}')[1] in TBLPR_ORDER else 99
    for e in kids: tblPr.remove(e)
    for e in sorted(kids, key=key): tblPr.append(e)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true')
    trPr.append(h)

# marcaj în text: **aldin**, //italic//
TOK = re.compile(r'(\*\*.+?\*\*|//.+?//)', re.S)
def runs(p, text, size=None, color=None, italic=False, bold=False):
    for part in TOK.split(text):
        if not part: continue
        b, i = bold, italic
        if part.startswith('**') and part.endswith('**'): part, b = part[2:-2], True
        elif part.startswith('//') and part.endswith('//'): part, i = part[2:-2], True
        r = p.add_run(part); r.bold = b; r.italic = i
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color
    return p

def TITLE(t): doc.add_paragraph(t, style='Title')
def H1(t): doc.add_paragraph(t, style='Heading 1')
def H2(t): doc.add_paragraph(t, style='Heading 2')
def H3(t): doc.add_paragraph(t, style='Heading 3')

def P(t):
    p = doc.add_paragraph(); runs(p, t)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def SMALL(t):
    p = doc.add_paragraph(); runs(p, t, size=9, color=GREY, italic=True)
    p.paragraph_format.space_after = Pt(8)
    return p

def QUOTE(t):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.right_indent = Cm(0.6)
    runs(p, t, size=10, color=BROWN, italic=True)
    return p

def BUL(items):
    for it in items:
        p = doc.add_paragraph(style='List Bullet'); runs(p, it)
        p.paragraph_format.space_after = Pt(3)

def NUM(items):
    for k, it in enumerate(items, 1):
        p = doc.add_paragraph(); runs(p, '**%d.** ' % k + it)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(3)

def TBL(rows, widths=None, header=True, small=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    t.autofit = False
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci)
            c.text = ''
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
            first = True
            for line in str(val).split('\n'):
                tgt = p if first else c.add_paragraph()
                if not first:
                    tgt.paragraph_format.space_before = Pt(0); tgt.paragraph_format.space_after = Pt(2)
                runs(tgt, line, size=8.5 if small else 9.5,
                     bold=(header and ri == 0), color=GREEN if (header and ri == 0) else None)
                first = False
            if header and ri == 0: shade(c, HEAD_FILL)
    if header: repeat_header(t.rows[0])
    if widths:
        set_widths(t, widths)
    order_tblPr(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def set_widths(t, widths):
    """Lățimi în cm, scrise atât în celule (Word) cât și în tblGrid (LibreOffice)."""
    for row in t.rows:
        for ci, w in enumerate(widths):
            row.cells[ci].width = Cm(w)
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(w * 567)))
    tblPr = t._tbl.tblPr
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(widths) * 567))); tw.set(qn('w:type'), 'dxa')
    for old in tblPr.findall(qn('w:tblW')): tblPr.remove(old)
    tblPr.append(tw)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)

def CALLOUT(title, body, fill='FFF6E5', line='E8C97A'):
    """body: text sau listă de rânduri (fiecare devine paragraf separat)."""
    tbl = doc.add_table(rows=1, cols=1); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_widths(tbl, [17.0])
    c = tbl.cell(0, 0); shade(c, fill); borders(tbl, line, 6)
    p = c.paragraphs[0]; runs(p, title, size=10, bold=True)
    p.paragraph_format.space_before = Pt(3)
    for b in ([body] if isinstance(body, str) else body):
        q = c.add_paragraph(); runs(q, b, size=9.5)
        q.paragraph_format.space_after = Pt(3)
    order_tblPr(tbl)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def IMG(path, width_cm, caption=None):
    doc.add_picture(path, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        p = doc.add_paragraph(); runs(p, caption, size=8.5, color=GREY, italic=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

def PAGEBREAK():
    doc.add_page_break()

def FOOTER(text):
    f = doc.sections[0].footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(f, text + '  ·  pagina ', size=8, color=GREY)
    r = f.add_run(); r.font.size = Pt(8); r.font.color.rgb = GREY
    for kind, val in (('begin', None), ('instr', 'PAGE'), ('end', None)):
        if kind == 'instr':
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = val
            r._r.append(it)
        else:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind)
            r._r.append(fc)
