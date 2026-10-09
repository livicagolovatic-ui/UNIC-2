# -*- coding: utf-8 -*-
"""Extrage rândurile unui expert din „Anexa 12 – Planificare lunară activități” (.docx).

În Anexa 12, coloanele „Activitate” și „Persoană și date de contact” sunt celule îmbinate vertical:
valoarea scrisă o dată se aplică rândurilor de sub ea până la următoarea valoare. Scriptul rezolvă
îmbinările și listează sesiunile, apoi semnalează automat:
  - sesiuni mai scurte de 60 de minute;
  - intervale identice / suprapuse cu alți experți, la aceeași unitate;
  - data documentului ulterioară primei activități planificate.

Utilizare:
  python3 extrage_anexa12.py Anexa_12.docx                       # implicit: „Facilitator comunitar 2”
  python3 extrage_anexa12.py Anexa_12.docx --persoana "Expert comunicare" --json iesire.json
"""
import sys, re, json, argparse
from datetime import datetime
import docx
from docx.oxml.ns import qn

ap = argparse.ArgumentParser()
ap.add_argument('anexa')
ap.add_argument('--persoana', default='Facilitator comunitar 2')
ap.add_argument('--json')
args = ap.parse_args()

d = docx.Document(args.anexa)
text_doc = '\n'.join(p.text for p in d.paragraphs)
tbl = max(d.tables, key=lambda t: len(t.rows))._tbl


def ctext(c):
    return ' '.join(''.join(x.text or '' for x in p.iter(qn('w:t'))) for p in c.findall(qn('w:p'))).strip()


rows, carry = [], {}
trs = tbl.findall(qn('w:tr'))
header = [ctext(c) for c in trs[0].findall(qn('w:tc'))]
for tr in trs[1:]:
    cells = tr.findall(qn('w:tc'))
    vals = []
    for j, c in enumerate(cells):
        vm = c.find('.//' + qn('w:vMerge'))
        t = ctext(c)
        if vm is not None and vm.get(qn('w:val')) != 'restart' and not t:
            t = carry.get(j, '')
        carry[j] = t
        vals.append(t)
    if len(vals) >= 8:
        rows.append(dict(activitate=vals[0], locatie=re.sub(r'\s+', ' ', vals[1]), modalitate=vals[2], data=vals[3].strip(),
                         interval=re.sub(r'\s+', '', vals[4]).replace('.', ':').replace('–', '-'), persoana=vals[6],
                         gt=re.sub(r'\s+', ' ', vals[7]).strip(), mentiuni=vals[8] if len(vals) > 8 else ''))


def minute(iv):
    m = re.match(r'(\d{1,2}):(\d{2})-(\d{1,2}):(\d{2})', iv)
    if not m:
        return None, None, None
    a, b = int(m[1]) * 60 + int(m[2]), int(m[3]) * 60 + int(m[4])
    return a, b, b - a


mine = [r for r in rows if args.persoana.lower() in r['persoana'].lower()]
alerts = []
for r in mine:
    a, b, dur = minute(r['interval'])
    r['durata_min'] = dur
    if dur is not None and dur < 60:
        alerts.append('%s %s – sesiune de %d min (materialele standard sunt de 60 min).' % (r['data'], r['interval'], dur))
    for o in rows:
        if o is r or args.persoana.lower() in o['persoana'].lower() or o['data'] != r['data']:
            continue
        oa, ob, _ = minute(o['interval'])
        if a is None or oa is None:
            continue
        loc_r, loc_o = r['locatie'].lower(), o['locatie'].lower()
        same_unit = loc_r[:25] == loc_o[:25] or ('turda' in loc_r and 'turda' in loc_o) or ('huedin' in loc_r and 'huedin' in loc_o)
        if same_unit and a < ob and oa < b:
            alerts.append('%s %s – se suprapune cu „%s” (%s, %s) la aceeași unitate: risc de dublă raportare a orelor / GT.'
                          % (r['data'], r['interval'], o['persoana'].split('tel')[0].strip(), o['interval'], o['activitate'][:6]))
m = re.search(r'Nr\.?\s*([\d/]+)\s*/\s*(\d{2}\.\d{2}\.\d{4})', text_doc)
if m and mine:
    try:
        ddoc = datetime.strptime(m[2], '%d.%m.%Y')
        first = min(datetime.strptime(r['data'], '%d.%m.%Y') for r in mine)
        if ddoc > first:
            alerts.append('Documentul este datat %s, după prima activitate planificată (%s) – verificați data.' % (m[2], first.strftime('%d.%m.%Y')))
    except ValueError:
        pass

print('Sesiuni pentru „%s”: %d (total %s min)' % (args.persoana, len(mine), sum(r['durata_min'] or 0 for r in mine)))
for r in mine:
    print(' - %s | %s | %s min | %s | %s | GT: %s' % (r['data'], r['interval'], r['durata_min'], r['activitate'][:6],
                                                      r['locatie'][:70], r['gt']))
print('\nDe verificat:' if alerts else '\nNicio neconcordanță detectată automat.')
for x in dict.fromkeys(alerts):
    print(' ! ' + x)
if args.json:
    json.dump({'persoana': args.persoana, 'sesiuni': mine, 'alerte': list(dict.fromkeys(alerts))},
              open(args.json, 'w'), ensure_ascii=False, indent=2)
