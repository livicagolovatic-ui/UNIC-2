# -*- coding: utf-8 -*-
import math, sys, copy
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
sys.path.insert(0, '/home/user/UNIC-2/annex')
import timetable_content as C

SRC = '/home/user/UNIC-2/annex/KA152_Annex_Timetable_TEMPLATE.xlsx'
OUT = '/home/user/UNIC-2/VERIFAI_KA152_Annex_Timetable.xlsx'

wb = load_workbook(SRC)

# ------------------------------------------------------------------ helpers
def unmerge_if(ws, ref):
    if ref in {str(m) for m in ws.merged_cells.ranges}:
        ws.unmerge_cells(ref)

def merge_if(ws, ref):
    if ref not in {str(m) for m in ws.merged_cells.ranges}:
        ws.merge_cells(ref)

def stamp(ws, row, tmpl_row, cols='ABCDEFGHIJ'):
    """Copy the cell styles of tmpl_row onto row."""
    for col in cols:
        src = ws['%s%d' % (col, tmpl_row)]
        dst = ws['%s%d' % (col, row)]
        dst._style = copy.copy(src._style)

def put(ws, coord, text, size=10, bold=False, halign='left', valign='top', wrap=True):
    c = ws[coord]
    c.value = text
    f = c.font
    c.font = Font(name='Arial', sz=size, bold=bold, color=f.color, italic=f.italic)
    c.alignment = Alignment(horizontal=halign, vertical=valign, wrap_text=wrap)
    return c

def height_for(texts, widths, size=10):
    """Rough row height: widest requirement of the cells sharing the row."""
    need = 1
    for t, w in zip(texts, widths):
        if not t:
            continue
        lines = 0
        for para in str(t).split('\n'):
            lines += max(1, math.ceil(len(para) / w))
        need = max(need, lines)
    return max(18.0, need * (size * 1.32) + 4)

# =================================================================== SHEET 1
ws = wb['Youth Exchanges']
W_ACT, W_MET = 60, 68          # usable chars in merged B:E and F:J

put(ws, 'B2', '01', size=11, halign='center', valign='center', wrap=False)
put(ws, 'B3', C.ORGS, size=10)
ws.row_dimensions[3].height = height_for([C.ORGS], [150])
put(ws, 'B4', C.DURATION, size=10)
ws.row_dimensions[4].height = height_for([C.DURATION], [150])
put(ws, 'A7', C.CITY,    size=10, halign='center', valign='center', wrap=False)
put(ws, 'D7', C.COUNTRY, size=10, halign='center', valign='center', wrap=False)
put(ws, 'G7', C.START,   size=10, halign='center', valign='center', wrap=False)
put(ws, 'I7', C.END,     size=10, halign='center', valign='center', wrap=False)

FOOT = 'For additional days, please copy the above rows'
unmerge_if(ws, 'A25:J25')
ws['A25'].value = None

BLOCK = 5                       # template block: day header + AM + AM cont + PM + PM cont
r = 10
for i, (label, am, pm) in enumerate(C.DAYS):
    # --- day header row
    stamp(ws, r, 10)
    unmerge_if(ws, 'B%d:E%d' % (r, r)); unmerge_if(ws, 'F%d:J%d' % (r, r))
    merge_if(ws, 'A%d:J%d' % (r, r))
    put(ws, 'A%d' % r, label, size=10, bold=True, halign='center', valign='center')
    ws.row_dimensions[r].height = height_for([label], [145])
    r += 1
    # --- content rows: AM slots then PM slots
    slots = [('AM' if k == 0 else '', t) for k, t in enumerate(am)] + \
            [('PM' if k == 0 else '', t) for k, t in enumerate(pm)]
    for j, (tag, (act, met)) in enumerate(slots):
        stamp(ws, r, 11 + min(j, 3))
        unmerge_if(ws, 'A%d:J%d' % (r, r))
        merge_if(ws, 'B%d:E%d' % (r, r)); merge_if(ws, 'F%d:J%d' % (r, r))
        put(ws, 'A%d' % r, tag or None, size=10, bold=True, halign='center', valign='center', wrap=False)
        put(ws, 'B%d' % r, act or None, size=10)
        put(ws, 'F%d' % r, met or None, size=10)
        ws.row_dimensions[r].height = height_for([act, met], [W_ACT, W_MET])
        r += 1

foot_row = r
stamp(ws, foot_row, 25)
unmerge_if(ws, 'B%d:E%d' % (foot_row, foot_row)); unmerge_if(ws, 'F%d:J%d' % (foot_row, foot_row))
merge_if(ws, 'A%d:J%d' % (foot_row, foot_row))
put(ws, 'A%d' % foot_row, FOOT, size=11, bold=True, halign='center', valign='center', wrap=False)
ws.row_dimensions[foot_row].height = 15.0
for extra in range(foot_row + 1, 60):          # clear anything left below
    for col in 'ABCDEFGHIJ':
        cc = ws['%s%d' % (col, extra)]
        if not hasattr(cc, 'value'):
            continue
        cc.value = None

# =================================================================== SHEET 2
pv = wb['Preparatory Visits ']
W_PV = 130

put(pv, 'B2', '02', size=11, halign='center', valign='center', wrap=False)
put(pv, 'B3', C.PV_LINK, size=10, valign='center')
pv.row_dimensions[3].height = height_for([C.PV_LINK], [W_PV])
put(pv, 'B4', C.PV_ORGS, size=10)
pv.row_dimensions[4].height = height_for([C.PV_ORGS], [W_PV])
put(pv, 'A7', 'Beliș (Cluj County)', size=10, halign='center', valign='center', wrap=False)
put(pv, 'D7', 'Romania',            size=10, halign='center', valign='center', wrap=False)
put(pv, 'G7', '16/06/2027',         size=10, halign='center', valign='center', wrap=False)
put(pv, 'I7', '17/06/2027',         size=10, halign='center', valign='center', wrap=False)
pv['A4'].alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)

PV_FOOT = pv['A19'].value
unmerge_if(pv, 'A19:J19'); pv['A19'].value = None

for i, (label, am, pm) in enumerate(C.PV_DAYS):
    top = 10 + i * BLOCK
    stamp(pv, top, 10)
    unmerge_if(pv, 'B%d:J%d' % (top, top))
    merge_if(pv, 'A%d:J%d' % (top, top))
    put(pv, 'A%d' % top, label, size=10, bold=True, halign='center', valign='center')
    pv.row_dimensions[top].height = height_for([label], [W_PV])
    for j, (tag, act) in enumerate([('AM', am[0]), ('', am[1]), ('PM', pm[0]), ('', pm[1])]):
        r = top + 1 + j
        stamp(pv, r, 11 + j)
        unmerge_if(pv, 'A%d:J%d' % (r, r))
        merge_if(pv, 'B%d:J%d' % (r, r))
        put(pv, 'A%d' % r, tag or None, size=10, bold=True, halign='center', valign='center', wrap=False)
        put(pv, 'B%d' % r, act, size=10)
        pv.row_dimensions[r].height = height_for([act], [W_PV])

pv_foot = 10 + len(C.PV_DAYS) * BLOCK
stamp(pv, pv_foot, 10)
for col in 'ABCDEFGHIJ':
    pv['%s%d' % (col, pv_foot)]._style = copy.copy(pv['A19']._style)
unmerge_if(pv, 'A%d:J%d' % (pv_foot, pv_foot))
merge_if(pv, 'A%d:J%d' % (pv_foot, pv_foot))
put(pv, 'A%d' % pv_foot, PV_FOOT, size=11, halign='center', valign='center', wrap=False)
pv.row_dimensions[pv_foot].height = 15.0

# ------------------------------------------------------------------ print setup
for sheet in (ws, pv):
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = '1:9'

wb.save(OUT)
print('saved', OUT)
print('YE rows used: 10 ..', foot_row, '| PV rows used: 10 ..', pv_foot)
