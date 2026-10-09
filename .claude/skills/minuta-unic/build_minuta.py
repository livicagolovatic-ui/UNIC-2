#!/usr/bin/env python3
"""Generează minuta unei întâlniri de management pentru proiectele UNIC (PIDS) și UNIC2 (PEO).

Reproduce șablonul minutelor existente în Drive (antet cu logo UE / Guvern / UNIC și
textul proiectului, subsol cu logo GAL, Times New Roman 12, structura secțiunilor).

Utilizare:
    python3 build_minuta.py minuta.json                 # -> .docx
    python3 build_minuta.py minuta.json --pdf --preview # + PDF și PNG pe pagini (verificare vizuală)

Conținutul se dă într-un fișier JSON; vezi examples/*.json și SKILL.md.
"""
import argparse
import json
import os
import re
import subprocess
import sys
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor, Emu

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")

PURPLE = RGBColor(0x70, 0x30, 0xA0)   # titluri/subtitluri PEO (#7030A0)
BLUE_LINE = "4472C4"                  # linia de sub titlul PEO
GREY_TXT = RGBColor(0x1E, 0x1E, 0x1E) # textul din antet
GREY_BORDER = "D9D9D9"                # chenarul tabelului de semnături PEO

# ---------------------------------------------------------------------------
# Datele fixe ale celor două proiecte (preluate din minutele din Drive)
# ---------------------------------------------------------------------------
PROIECTE = {
    "UNIC": {  # PIDS – folderul „PIDS- UNIC”
        "program": "PIDS",
        "antet_titlu": "Inițiativa Unitară pentru Familii și Copii în Napoca Porolissum (UNIC - Porolissum)",
        "antet_rand2": "Proiect cofinanțat din Fondul Social European Plus prin Programul Incluziune și "
                       "Demnitate Socială 2021-2027 | Cod MySMIS 329335",
        "titlu": "Minuta întâlnirii de lucru",
        "activitate": ["Activitatea 5- MANAGEMENTUL PROIECTULUI", "S.A.5.1 Management al proiectului"],
        "stil": "pids",
    },
    "UNIC2": {  # PEO – folderul „PEO-UNIC 2”
        "program": "PEO",
        "antet_titlu": "UNIC - Uniți pentru Nevoile Incluzive și Continuitatea Educației Elevilor",
        "antet_rand2": "Proiect cofinanțat din Fondul Social European Plus prin Programul Educație și "
                       "Ocupare 2021-2027 | Cod MySMIS 352704",
        "titlu": "Minuta întâlnirii de management",
        "identificare": [
            ("Beneficiar:", "Asociația Grupul de Acțiune Locală Napoca Porolissum"),
            ("Proiect:", "UNIC – „Uniți pentru Nevoile Incluzive și Continuitatea Educației Elevilor”"),
            ("Cod SMIS:", "352704"),
            ("Activitatea:", "A.1 - Management proiect, implementare monitorizare şi raportare"),
            ("Subactivitatea:", "S.A. 1.1- Management de proiect, monitorizare și raportare"),
        ],
        "stil": "peo",
    },
}
ALIAS = {"UNIC": "UNIC", "PIDS": "UNIC", "UNIC1": "UNIC", "UNICPIDS": "UNIC", "UNICPOROLISSUM": "UNIC",
         "UNIC2": "UNIC2", "PEO": "UNIC2", "UNIC2PEO": "UNIC2"}


def resolve_project(name):
    key = ALIAS.get(re.sub(r"[^A-Z0-9]", "", str(name).upper()))
    if key is None:
        sys.exit(f"Proiect necunoscut: {name!r} (folosește „UNIC” = PIDS sau „UNIC2” = PEO)")
    return key


# ---------------------------------------------------------------------------
# Utilitare docx
# ---------------------------------------------------------------------------
def set_run_font(run, name="Times New Roman", size=12, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(att), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def para(container, text="", *, bold=False, italic=False, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         before=0, after=6, color=None, font="Times New Roman", keep_next=False, indent=None,
         first_line=None, line=1.0):
    p = container.add_paragraph()
    pf = p.paragraph_format
    p.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    if indent is not None:
        pf.left_indent = Pt(indent)
    if first_line is not None:
        pf.first_line_indent = Pt(first_line)
    if text:
        set_run_font(p.add_run(text), font, size, bold, italic, color)
    return p


def label_value(container, label, value, *, size=12, after=4, value_bold=False):
    p = para(container, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=after)
    set_run_font(p.add_run(label), size=size, bold=True)
    if value:
        set_run_font(p.add_run(" " + value), size=size, bold=value_bold)
    return p


PPR_AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                  "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
                  "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                  "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                  "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def bottom_border(p, color, size_eighths=6, space=4):
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size_eighths))
    b.set(qn("w:space"), str(space))
    b.set(qn("w:color"), color)
    bdr.append(b)
    # Word cere ordinea din schemă (pBdr înaintea spacing/ind/jc), altfel „conținut ilizibil”
    ppr.insert_element_before(bdr, *PPR_AFTER_PBDR)


def table_borders(table, color, size_eighths, inside=True):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        if color is None or (not inside and edge.startswith("inside")):
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(size_eighths))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), color)
        borders.append(e)
    tblPr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                                "w:tblCaption", "w:tblDescription")


def table_fixed_widths(table, widths):
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.insert_element_before(layout, "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.twips if hasattr(w, "twips") else Emu(w).twips)))


def row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trPr.append(cs)


def keep_table_together(table):
    """Word mută tabelul întreg pe pagina următoare dacă nu încape (keep-with-next pe toate rândurile
    în afară de ultimul) – ca în minutele originale, unde semnăturile nu sunt rupte între pagini."""
    for row in table.rows[:-1]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    trPr.append(h)


def set_language(doc, lang="ro-RO"):
    rpr = doc.styles["Normal"].element.get_or_add_rPr()
    l = OxmlElement("w:lang")
    l.set(qn("w:val"), lang)
    l.set(qn("w:eastAsia"), lang)
    l.set(qn("w:bidi"), lang)
    rpr.append(l)


# ---------------------------------------------------------------------------
# Antet / subsol (identice ca geometrie la ambele proiecte, în versiunea din sept.–oct. 2026)
# ---------------------------------------------------------------------------
def build_header_footer(section, proj):
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.5)
    section.top_margin = Cm(4.7)
    section.bottom_margin = Cm(2.8)
    section.header_distance = Cm(1.3)
    section.footer_distance = Cm(1.2)

    header = section.header
    header.is_linked_to_previous = False
    usable = section.page_width - section.left_margin - section.right_margin
    widths = [Cm(8.0), Cm(5.0), usable - Cm(13.0)]
    t = header.add_table(1, 3, usable)
    table_borders(t, None, 0)
    table_fixed_widths(t, widths)
    logos = [("logo_ue.png", Pt(175.8), WD_ALIGN_PARAGRAPH.LEFT),
             ("logo_guvern.png", Pt(55.1), WD_ALIGN_PARAGRAPH.CENTER),
             ("logo_unic.png", Pt(83.9), WD_ALIGN_PARAGRAPH.RIGHT)]
    for cell, (img, w, al) in zip(t.rows[0].cells, logos):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = al
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(os.path.join(ASSETS, img), width=w)
    # antetul nou are deja un paragraf gol înaintea tabelului – îl eliminăm
    first = header.paragraphs[0]
    if not first.text and first._p.getnext() is not None:
        first._p.getparent().remove(first._p)

    p1 = header.add_paragraph()
    p1.paragraph_format.space_before = Pt(10)
    p1.paragraph_format.space_after = Pt(0)
    p1.paragraph_format.line_spacing = 1.0
    set_run_font(p1.add_run("Titlul proiectului: "), "Arial", 8.5, True, color=GREY_TXT)
    set_run_font(p1.add_run(proj["antet_titlu"]), "Arial", 8.5, False, color=GREY_TXT)
    p2 = header.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.0
    set_run_font(p2.add_run(proj["antet_rand2"]), "Arial", 8.5, False, color=GREY_TXT)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run().add_picture(os.path.join(ASSETS, "logo_gal.png"), width=Pt(74.2))


# ---------------------------------------------------------------------------
# Formatarea duratei, după stilul fiecărui proiect
# ---------------------------------------------------------------------------
def _minutes(a, b):
    t1, t2 = datetime.strptime(a, "%H:%M"), datetime.strptime(b, "%H:%M")
    return int((t2 - t1).total_seconds() // 60)


def _ro_count(n, sing, plural):
    if n == 1:
        return f"1 {sing}"
    return f"{n} de {plural}" if n % 100 >= 20 or n % 100 == 0 and n else f"{n} {plural}"


def durata_text(spec, stil):
    if spec.get("durata"):
        return spec["durata"]
    a, b = spec["ora_inceput"], spec["ora_sfarsit"]
    m = _minutes(a, b)
    if stil == "pids":
        # ex. „40 minute (12:00-12:40)”, „60 minute (15:00-16:00)”
        return f"{m} minute ({a}-{b})"
    h, r = divmod(m, 60)
    parts = []
    if h:
        parts.append("1 oră" if h == 1 else f"{h} ore")
    if r:
        parts.append(_ro_count(r, "minut", "minute"))
    dur = " și ".join(parts)
    # ex. „12:30–14:00, 1 oră și 30 de minute” / „10:00–10:30 (30 de minute)”
    return f"{a}–{b}, {dur}" if h else f"{a}–{b} ({dur})"


# ---------------------------------------------------------------------------
# Corpul minutei
# ---------------------------------------------------------------------------
def add_rich(p, text, size=12, color=None):
    """Permite **bold** în text (ex. nume de secțiuni sau termene)."""
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if chunk:
            set_run_font(p.add_run(chunk), size=size, bold=(i % 2 == 1), color=color)


def body_par(doc, text, after=6):
    p = para(doc, after=after)
    add_rich(p, text)
    return p


def build_pids(doc, spec, proj):
    if spec.get("afiseaza_activitatea", spec.get("tip") != "financiar"):
        for i, line in enumerate(proj["activitate"]):
            para(doc, line, italic=True, size=10, align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
        para(doc, after=0)
    para(doc, spec.get("titlu", proj["titlu"]), bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER,
         before=18, after=14)
    label_value(doc, "Data:", spec["data"])
    label_value(doc, "Durata:", durata_text(spec, "pids"))
    label_value(doc, "Modalitate de desfășurare:", spec["modalitate"], after=8)

    para(doc, "Scop:", bold=True, after=3, keep_next=True)
    body_par(doc, spec["scop"], after=8)

    para(doc, "Participanți:", bold=True, after=3, keep_next=True)
    for i, name in enumerate(spec["participanti"]):
        p = para(doc, after=2 if i < len(spec["participanti"]) - 1 else 10, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_run_font(p.add_run("- " + name))
    if spec.get("nota_participanti"):
        body_par(doc, spec["nota_participanti"], after=8)

    para(doc, "Aspecte discutate în cadrul întâlnirii:", bold=True, after=6, before=4, keep_next=True)
    if spec.get("introducere"):
        body_par(doc, spec["introducere"], after=10)
    for i, pt in enumerate(spec["puncte"], 1):
        para(doc, f"{i}. {pt['titlu']}", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=3,
             keep_next=True)
        for t in pt.get("paragrafe", []):
            body_par(doc, t)
        for m in pt.get("liste", []):
            p = para(doc, after=2)
            add_rich(p, "- " + m)

    c = spec.get("concluzii")
    if c:
        para(doc, c.get("titlu", "Concluzii și măsuri stabilite:"), bold=True, before=8, after=4,
             align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
        for t in c.get("paragrafe", []):
            body_par(doc, t)
        masuri = c.get("masuri", [])
        for i, m in enumerate(masuri):
            p = para(doc, after=3)
            add_rich(p, "- " + m.rstrip(";.") + (";" if i < len(masuri) - 1 else "."))

    if spec.get("semnaturi", True):
        para(doc, after=0)
        para(doc, "Semnătură participanți:", bold=True, before=18, after=10, align=WD_ALIGN_PARAGRAPH.LEFT,
             keep_next=True)
        usable = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
        semn = spec.get("semnatari", spec["participanti"])
        t = doc.add_table(rows=len(semn), cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        table_borders(t, "000000", 8)
        table_fixed_widths(t, [usable // 2, usable - usable // 2])
        for row, name in zip(t.rows, semn):
            row.height = Cm(1.45)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            row_cant_split(row)
            row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = row.cells[0].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(name), size=11)
        keep_table_together(t)


def build_peo(doc, spec, proj):
    for lab, val in spec.get("identificare", proj["identificare"]):
        p = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, after=1)
        set_run_font(p.add_run(lab), bold=True)
        set_run_font(p.add_run(" " + val))
    t = para(doc, spec.get("titlu", proj["titlu"]), bold=True, size=16, color=PURPLE,
             align=WD_ALIGN_PARAGRAPH.CENTER, before=28, after=12)
    bottom_border(t, BLUE_LINE, 6, 4)

    label_value(doc, "Data:", spec["data"])
    label_value(doc, "Interval orar și durată:", durata_text(spec, "peo"))
    label_value(doc, "Modalitate de desfășurare:", spec["modalitate"])
    p = para(doc, after=6)
    set_run_font(p.add_run("Scop:"), bold=True)
    add_rich(p, " " + spec["scop"])

    para(doc, "Participanți:", bold=True, after=3, keep_next=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    for i, name in enumerate(spec["participanti"]):
        p = para(doc, after=1 if i < len(spec["participanti"]) - 1 else 8, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_run_font(p.add_run("• " + name))
    if spec.get("nota_participanti"):
        body_par(doc, spec["nota_participanti"], after=8)
    if spec.get("introducere"):
        body_par(doc, spec["introducere"], after=8)

    para(doc, "Aspecte discutate", bold=True, color=PURPLE, before=4, after=6, keep_next=True,
         align=WD_ALIGN_PARAGRAPH.LEFT)
    puncte = list(spec["puncte"])
    c = spec.get("concluzii")
    if c:  # la PEO concluziile sunt ultimul punct numerotat
        puncte.append({"titlu": c.get("titlu", "Concluzii"), "paragrafe": c.get("paragrafe", []),
                       "liste": c.get("masuri", [])})
    for i, pt in enumerate(puncte, 1):
        para(doc, f"{i}. {pt['titlu']}", bold=True, color=PURPLE, align=WD_ALIGN_PARAGRAPH.LEFT,
             before=6, after=4, keep_next=True)
        for tx in pt.get("paragrafe", []):
            body_par(doc, tx)
        lst = pt.get("liste", [])
        for j, m in enumerate(lst):
            p = para(doc, after=2 if j < len(lst) - 1 else 6, indent=18, first_line=-12)
            add_rich(p, "•  " + m)

    if spec.get("semnaturi", True):
        para(doc, after=0)
        para(doc, "Semnăturile participanților:", bold=True, color=PURPLE, before=18, after=8,
             align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
        usable = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
        semn = spec.get("semnatari", spec["participanti"])
        t = doc.add_table(rows=len(semn) + 1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        table_borders(t, GREY_BORDER, 4)
        w0 = int(usable * 0.708)
        table_fixed_widths(t, [w0, usable - w0])
        hdr = t.rows[0]
        repeat_header(hdr)
        for cell, txt in zip(hdr.cells, ["Numele, prenumele și poziția", "Semnătura"]):
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            set_run_font(p.add_run(txt), size=11, bold=True, color=PURPLE)
        for row, name in zip(t.rows[1:], semn):
            row.height = Cm(2.3)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            row_cant_split(row)
            row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = row.cells[0].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(name), size=11)
        keep_table_together(t)


def add_photos(doc, photos, base_dir):
    if not photos:
        return
    import io
    from PIL import Image, ImageOps
    sec = doc.sections[0]
    max_w = sec.page_width - sec.left_margin - sec.right_margin - Cm(1.0)
    max_h = Cm(11.5)
    for i, ph in enumerate(photos):
        path = ph if os.path.isabs(ph) else os.path.join(base_dir, ph)
        # Normalizăm orice poză (JPEG cu profil ICC/CMYK, WEBP, PNG uriaș, rotație EXIF de telefon)
        # într-un JPEG RGB de max. 2000 px – python-docx refuză unele formate, iar pozele
        # originale umflă documentul (minuta din 13.08 avea 6 MB doar din 2 poze).
        im = ImageOps.exif_transpose(Image.open(path))
        if im.mode in ("RGBA", "LA", "P"):
            im = im.convert("RGBA")
            bg = Image.new("RGB", im.size, "white")
            bg.paste(im, mask=im.split()[-1])
            im = bg
        else:
            im = im.convert("RGB")
        im.thumbnail((2000, 2000))
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=88)
        buf.seek(0)
        w, h = im.size
        scale = min(max_w / w, max_h / h)
        p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=18 if i == 0 else 10, after=0)
        p.add_run().add_picture(buf, width=int(w * scale))


def build(spec, out_path, base_dir):
    key = resolve_project(spec.get("proiect"))
    proj = PROIECTE[key]

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(0)
    set_language(doc)
    doc.core_properties.title = f"{proj['titlu']} {spec['data']}"
    doc.core_properties.author = spec.get("autor", "Asociația GAL Napoca Porolissum")

    build_header_footer(doc.sections[0], proj)
    (build_pids if proj["stil"] == "pids" else build_peo)(doc, spec, proj)
    add_photos(doc, spec.get("poze", []), base_dir)
    doc.save(out_path)
    return key


def default_name(spec, key):
    mod = spec.get("sufix") or ("online" if "online" in spec["modalitate"].lower() else "fizic")
    if key == "UNIC":
        extra = "_financiar_UNIC" if spec.get("tip") == "financiar" else f"_{mod}"
        return f"Minuta_intalnirii_{spec['data']}{extra}.docx"
    return f"Minuta_UNIC2_{spec['data']}_{mod}.docx"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", help="fișierul JSON cu conținutul minutei")
    ap.add_argument("-o", "--out", help="fișierul .docx rezultat (implicit: numele uzual al minutei, "
                                        "în folderul ./minute/)")
    ap.add_argument("--pdf", action="store_true", help="exportă și PDF (LibreOffice)")
    ap.add_argument("--preview", action="store_true", help="generează PNG pentru fiecare pagină (implică --pdf)")
    a = ap.parse_args()

    with open(a.spec, encoding="utf-8") as f:
        spec = json.load(f)
    base_dir = os.path.dirname(os.path.abspath(a.spec))
    key = resolve_project(spec.get("proiect"))
    out = a.out or os.path.join("minute", default_name(spec, key))
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    key = build(spec, out, base_dir)
    print(f"[ok] {out}  (proiect {key} / {PROIECTE[key]['program']})")

    if a.pdf or a.preview:
        outdir = os.path.dirname(os.path.abspath(out))
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, out],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        pdf = os.path.splitext(os.path.abspath(out))[0] + ".pdf"
        print(f"[ok] {pdf}")
        if a.preview:
            prefix = os.path.splitext(pdf)[0] + "_pagina"
            subprocess.run(["pdftoppm", "-r", "70", "-png", pdf, prefix], check=True)
            pages = sorted(p for p in os.listdir(outdir)
                           if p.startswith(os.path.basename(prefix)) and p.endswith(".png"))
            for p in pages:
                print(f"[ok] {os.path.join(outdir, p)}")


if __name__ == "__main__":
    main()
