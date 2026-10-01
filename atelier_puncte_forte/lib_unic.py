# -*- coding: utf-8 -*-
"""Helper-e pentru generarea materialelor de atelier pe antetul proiectului UNIC."""
import os, re
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import _Cell

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, 'antet_UNIC.docx')

# Paleta preluată din logo-urile proiectului (UNIC mov, verde GAL, galben UNIC)
C = dict(purple='4A2A82', lav='F1ECF9', lav2='DCD0F0', green='3F7F22', lgreen='E7F2DE',
         yellow='F2B705', lyellow='FFF4D2', blue='1F4E9A', lblue='E4ECF8', grey='666666',
         lgrey='F2F2F2', ink='1C1C1C', orange='B9501A', lorange='FBE8DA', white='FFFFFF',
         cut='8C8C8C')

CONTENT_W = 17.0  # cm (A4 - 2 x 2 cm)
FONT = 'Calibri'


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- XML order helpers
TCPR_ORDER = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap',
              'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark']
TBLPR_ORDER = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize',
               'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders',
               'shd', 'tblLayout', 'tblCellMar', 'tblLook']
PPR_ORDER = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl',
             'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens', 'kinsoku',
             'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi',
             'adjustRightInd', 'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents',
             'suppressOverlap', 'jc', 'textDirection', 'textAlignment', 'textboxTightWrap',
             'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange']


def _insert_ordered(parent, el, order):
    name = el.tag.split('}')[1]
    old = parent.find(qn('w:' + name))
    if old is not None:
        parent.remove(old)
    idx = order.index(name)
    for child in parent:
        cname = child.tag.split('}')[1]
        if cname in order and order.index(cname) > idx:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def _border_el(tag, spec):
    e = OxmlElement('w:' + tag)
    if spec is None:
        e.set(qn('w:val'), 'nil')
    else:
        val, sz, color = spec
        e.set(qn('w:val'), val)
        e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color)
    return e


def cell_borders(cell, **edges):
    """edges: top/left/bottom/right = (val, sz, color) sau None (fără bordură)."""
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        if edge in edges:
            b.append(_border_el(edge, edges[edge]))
    _insert_ordered(tcPr, b, TCPR_ORDER)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), fill)
    _insert_ordered(tcPr, sh, TCPR_ORDER)


def cell_margins(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        if v is not None:
            e = OxmlElement('w:' + k)
            e.set(qn('w:w'), str(int(v * 567)))
            e.set(qn('w:type'), 'dxa')
            m.append(e)
    _insert_ordered(tcPr, m, TCPR_ORDER)


def table_borders(t, spec=None, inside=None):
    tblPr = t._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        b.append(_border_el(edge, spec))
    for edge in ('insideH', 'insideV'):
        b.append(_border_el(edge, inside))
    _insert_ordered(tblPr, b, TBLPR_ORDER)


def table_cell_margins(t, lr=0.15, tb=0.06):
    tblPr = t._tbl.tblPr
    m = OxmlElement('w:tblCellMar')
    for k, v in (('top', tb), ('left', lr), ('bottom', tb), ('right', lr)):
        e = OxmlElement('w:' + k)
        e.set(qn('w:w'), str(int(v * 567)))
        e.set(qn('w:type'), 'dxa')
        m.append(e)
    _insert_ordered(tblPr, m, TBLPR_ORDER)


def row_setup(row, height=None, exact=False, cant_split=True):
    if height is not None:
        row.height = Cm(height)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exact else WD_ROW_HEIGHT_RULE.AT_LEAST
    if cant_split:
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn('w:cantSplit')) is None:
            trPr.insert(0, OxmlElement('w:cantSplit'))


def table(c, rows, cols, widths, borders=None, inside=None, lr=0.15, tb=0.06, align='center'):
    t = c.add_table(rows=rows, cols=cols)
    t.autofit = False
    t.alignment = {'center': WD_TABLE_ALIGNMENT.CENTER, 'left': WD_TABLE_ALIGNMENT.LEFT}[align]
    tblPr = t._tbl.tblPr
    lay = OxmlElement('w:tblLayout')
    lay.set(qn('w:type'), 'fixed')
    _insert_ordered(tblPr, lay, TBLPR_ORDER)
    tw = OxmlElement('w:tblW')
    tw.set(qn('w:w'), str(int(sum(widths) * 567)))
    tw.set(qn('w:type'), 'dxa')
    _insert_ordered(tblPr, tw, TBLPR_ORDER)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(w * 567)))
    for r in t.rows:
        for cell, w in zip(r.cells, widths):
            cell.width = Cm(w)
    table_borders(t, borders, inside)
    table_cell_margins(t, lr, tb)
    return t


# ---------------------------------------------------------------- paragraphs
TOK = re.compile(r'(\*\*.+?\*\*|//.+?//)', re.S)


def runs(p, text, size=None, bold=False, italic=False, color=None, font=None):
    for part in TOK.split(text):
        if not part:
            continue
        b, i = bold, italic
        if part.startswith('**') and part.endswith('**'):
            part, b = part[2:-2], True
        elif part.startswith('//') and part.endswith('//'):
            part, i = part[2:-2], True
        r = p.add_run(part)
        r.bold = b
        r.italic = i
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = rgb(color)
        if font:
            r.font.name = font
    return p


def newp(c):
    """Paragraf nou; într-o celulă reutilizează paragraful gol de la final."""
    if isinstance(c, _Cell):
        kids = [k for k in c._tc if k.tag in (qn('w:p'), qn('w:tbl'))]
        last = kids[-1] if kids else None
        if last is not None and last.tag == qn('w:p') and not last.findall('.//' + qn('w:r')):
            from docx.text.paragraph import Paragraph
            return Paragraph(last, c)
    return c.add_paragraph()


def para(c, text='', size=None, bold=False, italic=False, color=None, align=None,
         before=0, after=4, keep=False, line=None, indent=None):
    p = newp(c)
    if text:
        runs(p, text, size, bold, italic, color)
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line
    if keep:
        pf.keep_with_next = True
    if indent is not None:
        pf.left_indent = Cm(indent)
    if align:
        p.alignment = {'center': WD_ALIGN_PARAGRAPH.CENTER, 'right': WD_ALIGN_PARAGRAPH.RIGHT,
                       'left': WD_ALIGN_PARAGRAPH.LEFT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    return p


def bullet(c, text, size=None, color=None, mark='▪', mark_color=None, indent=0.5, after=2, italic=False):
    p = newp(c)
    pf = p.paragraph_format
    pf.left_indent = Cm(indent)
    pf.first_line_indent = Cm(-0.42)
    pf.space_after = Pt(after)
    pf.space_before = Pt(0)
    pf.tab_stops.add_tab_stop(Cm(indent))
    r = p.add_run(mark + '\t')
    r.font.color.rgb = rgb(mark_color or C['purple'])
    r.bold = True
    if size:
        r.font.size = Pt(size)
    runs(p, text, size, False, italic, color)
    return p


def lines(c, n, width, size=11, before=7):
    """Rânduri de scris (tab cu linie de subliniere până la `width` cm)."""
    for _ in range(n):
        p = newp(c)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(width), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES)
        r = p.add_run('\t')
        r.font.size = Pt(size)
        r.font.color.rgb = rgb('9A9A9A')


def field(c, parts, size=11, before=7, after=0, bold_labels=True, color=None):
    """parts: [(eticheta, sfarsit_linie_cm), ...] pe același rând."""
    p = newp(c)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    for i, (label, end) in enumerate(parts):
        p.paragraph_format.tab_stops.add_tab_stop(Cm(end), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES)
        if label:
            r = p.add_run(('   ' if i else '') + label + ' ')
            r.bold = bold_labels
            r.font.size = Pt(size)
            if color:
                r.font.color.rgb = rgb(color)
        r = p.add_run('\t')
        r.font.size = Pt(size)
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def spacer(c, pt=4):
    p = newp(c)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run('')
    p.paragraph_format.line_spacing = Pt(pt)
    return p


def checkbox_grid(c, items, cols, width, size=10.5, box='☐'):
    rows = (len(items) + cols - 1) // cols
    t = table(c, rows, cols, [width / cols] * cols, lr=0.05, tb=0.04)
    for i, it in enumerate(items):
        cell = t.cell(i // cols, i % cols)
        p = para(cell, '', after=0)
        r = p.add_run(box + ' ')
        r.font.size = Pt(size + 2)
        r.font.color.rgb = rgb(C['purple'])
        runs(p, it, size)
    return t


# ---------------------------------------------------------------- document
class UnicDoc:
    def __init__(self, footer_text):
        self.d = docx.Document(TEMPLATE)
        body = self.d.element.body
        for el in list(body):
            if el.tag == qn('w:p'):
                body.remove(el)
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin, s.bottom_margin = Cm(3.6), Cm(1.8)
        s.header_distance, s.footer_distance = Cm(0.5), Cm(0.7)
        self._fix_header(s)
        self._footer(s, footer_text)
        self._styles()

    # antetul original are lățimea > A4: îl redimensionăm și îl ancorăm la pagină
    def _fix_header(self, s):
        hdr = s.header._element
        anchor = hdr.find('.//' + qn('wp:anchor'))
        W = 7200000
        H = int(1092200 * W / 7639050)
        ph = anchor.find(qn('wp:positionH'))
        ph.set('relativeFrom', 'page')
        ph.find(qn('wp:posOffset')).text = str((7560000 - W) // 2)
        pv = anchor.find(qn('wp:positionV'))
        pv.set('relativeFrom', 'page')
        pv.find(qn('wp:posOffset')).text = '190000'
        ext = anchor.find(qn('wp:extent'))
        ext.set('cx', str(W))
        ext.set('cy', str(H))
        for x in anchor.iter(qn('a:ext')):
            if x.get('cx'):
                x.set('cx', str(W))
                x.set('cy', str(H))
        wt = anchor.find(qn('wp:wrapTight'))
        if wt is not None:
            wt.addprevious(OxmlElement('wp:wrapNone'))
            anchor.remove(wt)

    def _footer(self, s, text):
        f = s.footer
        f.is_linked_to_previous = False
        p = f.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement('w:pBdr')
        bdr.append(_border_el('top', ('single', 6, C['lav2'])))
        _insert_ordered(pPr, bdr, PPR_ORDER)
        r = p.add_run(text + '   |   pag. ')
        r.font.size = Pt(8)
        r.font.color.rgb = rgb(C['grey'])
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), 'PAGE')
        rr = OxmlElement('w:r')
        rpr = OxmlElement('w:rPr')
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '16')
        col = OxmlElement('w:color')
        col.set(qn('w:val'), C['grey'])
        rpr.append(col)
        rpr.append(sz)
        rr.append(rpr)
        t = OxmlElement('w:t')
        t.text = '1'
        rr.append(t)
        fld.append(rr)
        p._p.append(fld)

    def _set_font(self, style, size=None, bold=None, color=None):
        rpr = style.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.insert(0, rf)
        for a in list(rf.attrib):
            del rf.attrib[a]
        for a in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
            rf.set(qn('w:' + a), FONT)
        if size:
            style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if color:
            style.font.color.rgb = rgb(color)

    def _styles(self):
        st = self.d.styles
        n = st['Normal']
        self._set_font(n, 10.5, color=C['ink'])
        n.paragraph_format.space_after = Pt(4)
        n.paragraph_format.line_spacing = 1.08
        for name, size, color in (('Heading 1', 15, C['purple']), ('Heading 2', 12.5, C['purple']),
                                  ('Heading 3', 11, C['ink'])):
            s = st[name]
            self._set_font(s, size, True, color)
            s.font.italic = False
            s.paragraph_format.space_before = Pt(12)
            s.paragraph_format.space_after = Pt(5)
            s.paragraph_format.keep_with_next = True
        for name in ('Header', 'Footer'):
            self._set_font(st[name])

    def save(self, path):
        self.d.save(path)

    # ------------------------------------------------------------ building blocks
    def banner(self, kicker, title, subtitle, theme_line):
        t = table(self.d, 1, 1, [CONTENT_W], lr=0.45, tb=0.25)
        cell = t.cell(0, 0)
        shade(cell, C['purple'])
        para(cell, kicker, 10, True, color=C['yellow'], after=2)
        para(cell, title, 21, True, color=C['white'], after=2, line=1.0)
        para(cell, subtitle, 10.5, color='E6DDF7', after=6)
        p = para(cell, '', after=2)
        r = p.add_run(' ' + theme_line + ' ')
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = rgb(C['ink'])
        rPr = r._r.get_or_add_rPr()
        sh = OxmlElement('w:shd')
        sh.set(qn('w:val'), 'clear')
        sh.set(qn('w:color'), 'auto')
        sh.set(qn('w:fill'), C['yellow'])
        rPr.append(sh)
        spacer(self.d, 6)

    def info_grid(self, pairs):
        """pairs: listă de (etichetă, valoare); 2 perechi pe rând."""
        rows = (len(pairs) + 1) // 2
        t = table(self.d, rows, 4, [2.9, 5.6, 2.9, 5.6], borders=('single', 4, C['lav2']),
                  inside=('single', 4, C['lav2']))
        for i, (k, v) in enumerate(pairs):
            r, cidx = i // 2, (i % 2) * 2
            kc, vc = t.cell(r, cidx), t.cell(r, cidx + 1)
            shade(kc, C['lav'])
            para(kc, k, 9.5, True, color=C['purple'], after=0)
            if v:
                para(vc, v, 9.5, after=0)
            else:
                lines(vc, 1, 5.2, 9.5, before=2)
        spacer(self.d, 4)

    def h(self, text, level=2):
        p = self.d.add_paragraph(style='Heading %d' % level)
        runs(p, text)
        if level <= 2:
            pPr = p._p.get_or_add_pPr()
            bdr = OxmlElement('w:pBdr')
            bdr.append(_border_el('bottom', ('single', 8, C['lav2'])))
            _insert_ordered(pPr, bdr, PPR_ORDER)
        return p

    def p(self, text, **kw):
        return para(self.d, text, **kw)

    def b(self, text, **kw):
        return bullet(self.d, text, **kw)

    def callout(self, label, items, fill, label_fill, label_color='FFFFFF', width=CONTENT_W):
        t = table(self.d, 1, 2, [2.7, width - 2.7], lr=0.2, tb=0.12)
        lc, cc = t.cell(0, 0), t.cell(0, 1)
        shade(lc, label_fill)
        shade(cc, fill)
        lc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        para(lc, label, 9.5, True, color=label_color, align='center', after=0)
        fill_items(cc, items)
        spacer(self.d, 6)
        return t

    def simple_table(self, header, rows, widths, size=9.5, head_fill=None, zebra=True, bold_first=False):
        t = table(self.d, len(rows) + 1, len(header), widths, borders=('single', 4, C['lav2']),
                  inside=('single', 4, C['lav2']))
        for j, htxt in enumerate(header):
            cell = t.cell(0, j)
            shade(cell, head_fill or C['purple'])
            para(cell, htxt, size, True, color=C['white'], after=0)
        row_setup(t.rows[0])
        trPr = t.rows[0]._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:tblHeader'))
        for i, r in enumerate(rows, start=1):
            row_setup(t.rows[i])
            for j, v in enumerate(r):
                cell = t.cell(i, j)
                if zebra and i % 2 == 0:
                    shade(cell, C['lgrey'])
                if isinstance(v, (list, tuple)):
                    fill_items(cell, v, size)
                else:
                    para(cell, v, size, bold=(bold_first and j == 0), after=0)
        spacer(self.d, 6)
        return t

    def activity(self, num, title, dur, interval, rows):
        """Antet de activitate + tabel cu blocuri FĂ / SPUNE / ÎNTREABĂ / MESAJ-CHEIE..."""
        t = table(self.d, 1, 3, [3.3, 10.2, 3.5], lr=0.2, tb=0.12)
        row_setup(t.rows[0])
        a, b_, c = t.cell(0, 0), t.cell(0, 1), t.cell(0, 2)
        shade(a, C['purple'])
        shade(b_, C['lav2'])
        shade(c, C['yellow'])
        for x in (a, b_, c):
            x.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        para(a, 'ACTIVITATEA %s' % num, 10, True, color=C['white'], align='center', after=0)
        para(b_, title, 13, True, color=C['purple'], after=0)
        para(c, dur, 11, True, color=C['ink'], align='center', after=0)
        para(c, interval, 8.5, color=C['ink'], align='center', after=0)
        t2 = table(self.d, len(rows), 2, [2.7, CONTENT_W - 2.7], borders=('single', 4, C['lav2']),
                   inside=('single', 4, C['lav2']), lr=0.2, tb=0.1)
        for i, (label, items) in enumerate(rows):
            lf, lc_, cf, ital = LABELS[label]
            l, cc = t2.cell(i, 0), t2.cell(i, 1)
            shade(l, lf)
            if cf:
                shade(cc, cf)
            para(l, label, 9.5, True, color=lc_, align='center', after=0, before=2)
            fill_items(cc, items, italic=ital)
        # primul rând al blocului rămâne cu antetul
        for pp in a.paragraphs + b_.paragraphs + c.paragraphs:
            pp.paragraph_format.keep_with_next = True
        spacer(self.d, 8)

    def annex_title(self, code, title, note=None, first=False):
        if not first:
            page_break(self.d)
        t = table(self.d, 1, 2, [3.0, CONTENT_W - 3.0], lr=0.25, tb=0.12)
        a, b_ = t.cell(0, 0), t.cell(0, 1)
        shade(a, C['purple'])
        shade(b_, C['lav'])
        for x in (a, b_):
            x.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        para(a, code, 11, True, color=C['white'], align='center', after=0)
        para(b_, title, 14, True, color=C['purple'], after=0)
        if note:
            para(self.d, note, 9, italic=True, color=C['grey'], before=3, after=4)
        else:
            spacer(self.d, 6)


# eticheta: (fundal etichetă, culoare text etichetă, fundal conținut, italic)
LABELS = {
    'SCOP': (C['lav2'], C['purple'], None, False),
    'PREGĂTIRE': (C['lblue'], C['blue'], None, False),
    'FĂ': (C['blue'], C['white'], None, False),
    'SPUNE': (C['purple'], C['white'], C['lav'], True),
    'ÎNTREABĂ': (C['green'], C['white'], C['lgreen'], False),
    'MESAJ-CHEIE': (C['yellow'], C['ink'], C['lyellow'], False),
    'ATENȚIE': (C['orange'], C['white'], C['lorange'], False),
    'VARIANTĂ': (C['grey'], C['white'], C['lgrey'], False),
    'MATERIALE': (C['lgrey'], C['ink'], None, False),
    'AFIRMAȚII': (C['lblue'], C['blue'], None, False),
    'PAȘI': (C['lblue'], C['blue'], None, False),
}


def fill_items(cell, items, size=10, italic=False):
    """Elemente: '- text' = bullet; '# text' = subtitlu; '> text' = replică; altfel paragraf."""
    if isinstance(items, str):
        items = [items]
    for it in items:
        if it.startswith('- '):
            bullet(cell, it[2:], size, italic=italic)
        elif it.startswith('# '):
            para(cell, it[2:], size, True, color=C['purple'], after=1, before=3, keep=True)
        elif it.startswith('> '):
            para(cell, it[2:], size, italic=True, color=C['purple'], after=3, indent=0.3)
        else:
            para(cell, it, size, italic=italic, after=3)


def cut_grid(doc, n, cols, w, h, fill, dashed=True):
    """Grilă de cartonașe de decupat; fill(cell, index)."""
    rows = (n + cols - 1) // cols
    spec = ('dashed', 6, C['cut']) if dashed else ('single', 6, C['cut'])
    t = table(doc, rows, cols, [w] * cols, borders=spec, inside=spec, lr=0.25, tb=0.15)
    for r in t.rows:
        row_setup(r, h, exact=True)
    for i in range(rows * cols):
        cell = t.cell(i // cols, i % cols)
        if i < n:
            fill(cell, i)
    return t


def scissors_note(doc, text='✂  Decupează pe linia punctată.'):
    para(doc, text, 8.5, italic=True, color=C['grey'], after=4)
