# -*- coding: utf-8 -*-
"""Planificarea Facilitatorului comunitar 2 pentru Huedin – octombrie 2026.

Utilizare:  python3 build_planificare_huedin.py <Repartizarea_copii_pe_clase.xlsx> <iesire.xlsx>

Numele elevilor se citesc din fișierul de repartizare primit de la școală și NU sunt
păstrate în acest script (date cu caracter personal ale unor minori).
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from datetime import date, time

SRC, OUT = sys.argv[1], sys.argv[2]

FONT = 'Arial'
PURPLE, LAV, YELLOW, GREY, GREEN, ORANGE = '4A2A82', 'EFE9F8', 'FFF2CC', 'F2F2F2', 'E2EFDA', 'FCE4D6'
thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def f(bold=False, color='000000', size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def fill(c):
    return PatternFill('solid', start_color=c, end_color=c)


def hdr(ws, row, headers, widths=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = f(True, 'FFFFFF')
        c.fill = fill(PURPLE)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = BORDER
    ws.row_dimensions[row].height = 32
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w


def cell(ws, row, col, value, bold=False, wrap=True, bg=None, num=None, align='left', color='000000', italic=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = f(bold, color, italic=italic)
    c.alignment = Alignment(wrap_text=wrap, vertical='top', horizontal=align)
    c.border = BORDER
    if bg:
        c.fill = fill(bg)
    if num:
        c.number_format = num
    return c


def title(ws, text, sub=None):
    ws['A1'] = text
    ws['A1'].font = f(True, PURPLE, 14)
    if sub:
        ws['A2'] = sub
        ws['A2'].font = f(False, '595959', 9, italic=True)


def page(ws, landscape=True):
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


# ------------------------------------------------------------------ citire elevi
src = openpyxl.load_workbook(SRC, data_only=True)
elevi = []  # (nume, clasa)
for ws in src.worksheets:
    clasa = None
    for row in ws.iter_rows(values_only=True):
        for v in row:
            if isinstance(v, str) and v.strip().upper().startswith('CLASA'):
                clasa = v.strip().upper().replace('CLASA A ', '').replace('CLASA ', '')
        if clasa and isinstance(row[0], (int, float)) and len(row) > 1 and row[1]:
            elevi.append((' '.join(str(row[1]).split()), clasa))
CLASE = []
for _, c in elevi:
    if c not in CLASE:
        CLASE.append(c)

GRUPA_CLASA = {'VII-A': 'Grupa 1', 'VII-B': 'Grupa 1', 'VIII-A': 'Grupa 2', 'VIII-B': 'Grupa 2'}

wb = openpyxl.Workbook()

# ================================================================== 1. Sinteză
ws = wb.active
ws.title = 'Sinteză'
title(ws, 'UNIC 2 – Huedin: planificarea Facilitatorului comunitar 2, octombrie 2026',
      'Celulele galbene se pot modifica (repartizarea claselor pe grupe, GT declarat, prezențe). Restul se calculează automat.')
info = [('Proiect', 'UNIC – Uniți pentru Nevoile Incluzive și Continuitatea Educației Elevilor, cod SMIS 352704'),
        ('Subactivitate', 'SA5.3 – Dezvoltarea de programe de informare și conștientizare'),
        ('Unitatea de învățământ', 'Liceul Tehnologic „Vlădeasa”, Piața Republicii nr. 39–42, Huedin, jud. Cluj'),
        ('Expert', 'Facilitator comunitar 2 – Golovatic Livia'),
        ('Program', 'Atelierul „Eu, punctele mele forte și ce mă motivează” (2 ședințe / grupă)'),
        ('Sursa programării', 'Anexa 12 – Planificare activități luna octombrie 2026, versiunea 2')]
for i, (k, v) in enumerate(info, start=4):
    cell(ws, i, 1, k, True, bg=LAV)
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=5)
    cell(ws, i, 2, v)
ws.column_dimensions['A'].width = 26
for col, w in zip('BCDE', (16, 16, 18, 30)):
    ws.column_dimensions[col].width = w

r0 = 11
ws.cell(row=r0 - 1, column=1, value='Repartizarea elevilor pe clase și grupe').font = f(True, PURPLE, 11)
hdr(ws, r0, ['Clasa', 'Grupa (se poate modifica)', 'Nr. elevi'])
CLS_FIRST = r0 + 1
for i, c in enumerate(CLASE):
    r = CLS_FIRST + i
    cell(ws, r, 1, c, True)
    cell(ws, r, 2, GRUPA_CLASA.get(c, 'Grupa 1'), bg=YELLOW, align='center')
    cell(ws, r, 3, "=COUNTIF('Elevi pe grupe'!$C:$C,A%d)" % r, align='center')
CLS_LAST = CLS_FIRST + len(CLASE) - 1
r = CLS_LAST + 1
cell(ws, r, 1, 'TOTAL', True, bg=GREY)
cell(ws, r, 2, '', bg=GREY)
cell(ws, r, 3, '=SUM(C%d:C%d)' % (CLS_FIRST, CLS_LAST), True, bg=GREY, align='center')
dv = DataValidation(type='list', formula1='"Grupa 1,Grupa 2"', allow_blank=False)
ws.add_data_validation(dv)
dv.add('B%d:B%d' % (CLS_FIRST, CLS_LAST))

g0 = r + 3
ws.cell(row=g0 - 1, column=1, value='Grupele de lucru').font = f(True, PURPLE, 11)
hdr(ws, g0, ['Grupa', 'Clase', 'Nr. elevi', 'Sesiuni în octombrie', 'Ore cu grupul țintă'])
GRP_FIRST = g0 + 1
for i, g in enumerate(('Grupa 1', 'Grupa 2')):
    rr = GRP_FIRST + i
    cell(ws, rr, 1, g, True)
    clase = [c for c in CLASE if GRUPA_CLASA.get(c) == g]
    cell(ws, rr, 2, ' + '.join(clase), color='000000')
    ws.cell(row=rr, column=2).comment = Comment('Propunere: clasele a VII-a în Grupa 1, clasele a VIII-a în Grupa 2. '
                                                'Dacă modificați repartizarea în tabelul de mai sus, actualizați și acest text.', 'UNIC')
    cell(ws, rr, 3, "=COUNTIF('Elevi pe grupe'!$D:$D,A%d)" % rr, align='center')
    cell(ws, rr, 4, "=COUNTIF('Planificare Huedin oct.'!$G:$G,A%d)" % rr, align='center')
    cell(ws, rr, 5, "=SUMIF('Planificare Huedin oct.'!$G:$G,A%d,'Planificare Huedin oct.'!$F:$F)/60" % rr, num='0.00', align='center')
GRP_LAST = GRP_FIRST + 1

k0 = GRP_LAST + 3
ws.cell(row=k0 - 1, column=1, value='Indicatori – octombrie 2026, Huedin').font = f(True, PURPLE, 11)
hdr(ws, k0, ['Indicator', 'Valoare'])
ind = [
    ('Sesiuni planificate la Huedin', "=COUNT('Planificare Huedin oct.'!B5:B8)", '0'),
    ('Ore cu grupul țintă la Huedin', "=SUM('Planificare Huedin oct.'!F5:F8)/60", '0.00'),
    ('Elevi repartizați (total)', '=C%d' % (CLS_LAST + 1), '0'),
    ('GT declarat în Anexa 12 (total pe sesiuni)', "=SUM('Planificare Huedin oct.'!J5:J8)", '0'),
    ('Elevi prezenți – Ședința 1', "=COUNTIF('Elevi pe grupe'!E:E,\"P\")", '0'),
    ('Elevi prezenți – Ședința 2', "=COUNTIF('Elevi pe grupe'!F:F,\"P\")", '0'),
    ('Elevi care au participat la ambele ședințe', "=COUNTIFS('Elevi pe grupe'!E:E,\"P\",'Elevi pe grupe'!F:F,\"P\")", '0'),
]
for i, (k, v, nf) in enumerate(ind, start=k0 + 1):
    cell(ws, i, 1, k, True, bg=LAV)
    cell(ws, i, 2, v, num=nf, align='center')
page(ws, landscape=False)

# ================================================================== 2. Planificare
wp = wb.create_sheet('Planificare Huedin oct.')
title(wp, 'Planificarea sesiunilor SA5.3 la Huedin – Facilitator comunitar 2 (octombrie 2026)',
      'Datele, intervalele orare și GT declarat sunt preluate din Anexa 12 (V2). Grupa, tematica și activitățile sunt propuse.')
H = ['Nr.', 'Data', 'Ziua', 'Ora început', 'Ora sfârșit', 'Durata (min)', 'Grupa', 'Clase', 'Nr. elevi în grupă',
     'GT declarat Anexa 12', 'Diferență (grupă – declarat)', 'Ședința', 'Tematica abordată', 'Activități propuse',
     'Materiale (anexe)', 'Elevi prezenți', 'Observații']
hdr(wp, 4, H, [5, 11, 9, 9, 9, 9, 9, 14, 9, 10, 11, 8, 30, 58, 26, 9, 40])

S1 = 'Ședința 1 – „Cine sunt eu și ce am bun?” (autocunoaștere, puncte forte, feedback pozitiv)'
S2 = 'Ședința 2 – „Ce mă motivează și ce vreau să dezvolt?” (motivație, superputere, obiectiv de 30 de zile)'
ACT_S1 = ('Recrutarea agenților (ecusoane, reguli); Schimbă locul dacă...; Dosarul secret – ghicește talentul; Detectivii de puncte forte '
          '(interviu în perechi); Harta echipei + Misiunea imposibilă; Mingea complimentelor; Cardul de agent + biletul de ieșire.')
ACT_S1_40 = ('VARIANTA SCURTĂ (40 min): Recrutarea agenților; Schimbă locul dacă... (scurt); Detectivii de puncte forte; Harta echipei '
             '(fără Misiunea imposibilă); Mingea complimentelor; Cardul de agent + biletul de ieșire. Se omite Dosarul secret.')
ACT_S2 = ('Reconectare – bateria mea; Colțurile motivației; Motivația se schimbă (combustibil și frâne); Superputerea mea; '
          'Provocarea LEVEL UP 30 + semnul de carte; Licitația super-echipei; Scrisoare către mine, peste un an; Încheiere + diplome + feedback.')
MAT_S1 = 'Ședința 1, Anexele 1–9 (ecusoane, bilețele, Fișele 1–2, indicatoare, cartonașe, Card de agent, bilet de ieșire)'
MAT_S2 = 'Ședința 2, Anexele 1–10 (cartonașe, Fișele 3–7, semn de carte, fișă de feedback, diplome) + plicuri'
plan = [
    (date(2026, 10, 9), time(10, 15), time(10, 55), 'Grupa 2', 2, 1, S1, ACT_S1_40, MAT_S1,
     'Interval de 40 min în Anexa 12 – se aplică varianta scurtă a Ședinței 1 (vezi foaia „Tematici și activități”).'),
    (date(2026, 10, 9), time(11, 55), time(12, 55), 'Grupa 1', 3, 1, S1, ACT_S1, MAT_S1, ''),
    (date(2026, 10, 29), time(11, 0), time(12, 0), 'Grupa 2', 3, 2, S2, ACT_S2, MAT_S2,
     'Același interval apare în Anexa 12 și la Expertul comunicare – de clarificat (vezi „Observații”).'),
    (date(2026, 10, 29), time(12, 0), time(13, 0), 'Grupa 1', 3, 2, S2, ACT_S2, MAT_S2,
     'Același interval apare în Anexa 12 și la Expertul comunicare – de clarificat (vezi „Observații”).'),
]
FIRST = 5
for i, (dt, t1, t2, grp, decl, sed, tem, act, mat, obs) in enumerate(plan):
    r = FIRST + i
    cell(wp, r, 1, i + 1, align='center')
    cell(wp, r, 2, dt, num='dd.mm.yyyy', align='center')
    cell(wp, r, 3, '=CHOOSE(WEEKDAY(B%d,2),"Luni","Marți","Miercuri","Joi","Vineri","Sâmbătă","Duminică")' % r, align='center')
    cell(wp, r, 4, t1, num='hh:mm', align='center')
    cell(wp, r, 5, t2, num='hh:mm', align='center')
    cell(wp, r, 6, '=ROUND((E%d-D%d)*1440,0)' % (r, r), align='center')
    cell(wp, r, 7, grp, True, bg=YELLOW, align='center')
    cell(wp, r, 8, "=IFERROR(INDEX('Sinteză'!$B$%d:$B$%d,MATCH(G%d,'Sinteză'!$A$%d:$A$%d,0)),\"\")"
         % (GRP_FIRST, GRP_LAST, r, GRP_FIRST, GRP_LAST), align='center')
    cell(wp, r, 9, "=COUNTIF('Elevi pe grupe'!$D:$D,G%d)" % r, align='center')
    cell(wp, r, 10, decl, bg=YELLOW, align='center')
    cell(wp, r, 11, '=I%d-J%d' % (r, r), align='center')
    cell(wp, r, 12, sed, align='center')
    cell(wp, r, 13, tem)
    cell(wp, r, 14, act)
    cell(wp, r, 15, mat)
    cell(wp, r, 16, "=IF(L%d=1,COUNTIFS('Elevi pe grupe'!$D:$D,G%d,'Elevi pe grupe'!$E:$E,\"P\"),"
                    "COUNTIFS('Elevi pe grupe'!$D:$D,G%d,'Elevi pe grupe'!$F:$F,\"P\"))" % (r, r, r), align='center')
    cell(wp, r, 17, obs, bg=ORANGE if obs else None)
    wp.row_dimensions[r].height = 92
LAST = FIRST + len(plan) - 1
r = LAST + 1
cell(wp, r, 1, '', bg=GREY)
cell(wp, r, 2, 'TOTAL', True, bg=GREY)
for col in (3, 4, 5, 7, 8, 12, 13, 14, 15, 17):
    cell(wp, r, col, '', bg=GREY)
cell(wp, r, 6, '=SUM(F%d:F%d)' % (FIRST, LAST), True, bg=GREY, align='center')
cell(wp, r, 9, '', bg=GREY)
cell(wp, r, 10, '=SUM(J%d:J%d)' % (FIRST, LAST), True, bg=GREY, align='center')
cell(wp, r, 11, '', bg=GREY)
cell(wp, r, 16, '=SUM(P%d:P%d)' % (FIRST, LAST), True, bg=GREY, align='center')
wp.cell(row=r + 2, column=1, value='Legendă: celulele galbene se pot modifica; „Diferență” pozitivă = mai mulți elevi în grupă decât GT declarat '
                                   'în Anexa 12; „Elevi prezenți” se completează automat din foaia „Elevi pe grupe” (P/A).').font = f(italic=True, size=9, color='595959')
wp.freeze_panes = 'C5'
page(wp)

# ================================================================== 3. Tematici
wt = wb.create_sheet('Tematici și activități')
title(wt, 'Tematicile abordate în octombrie și activitățile propuse (atelierul „Eu, punctele mele forte și ce mă motivează”)',
      'Minutul de început se calculează din durate; totalul fiecărui bloc trebuie să fie egal cu durata sesiunii.')
widths = [12, 9, 34, 70, 30]
for j, w in enumerate(widths, start=1):
    wt.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
blocks = [
    ('ȘEDINȚA 1 – „Cine sunt eu și ce am bun?” (60 min) – Grupa 1, 09.10.2026 | tema creativă: Agenția Secretă a Punctelor Forte',
     'Scop: autocunoaștere prin identificarea punctelor forte și a resurselor personale; exersarea feedbackului pozitiv.', [
        (5, '0. Recrutarea agenților', 'Primirea elevilor, ecusoane cu nume de cod, cele 3 reguli ale grupului.', 'Anexa 1'),
        (7, '1. Schimbă locul dacă...', 'Energizare: elevii observă că au preferințe și abilități diferite.', '–'),
        (8, '2. Dosarul secret – ghicește talentul', 'Fiecare scrie anonim un lucru la care se pricepe; grupul ghicește.', 'Anexa 2, cutie'),
        (12, '3. Detectivii de puncte forte', 'Interviu în perechi cu „dovezi”; raportul detectivului.', 'Anexa 3 (Fișa 1)'),
        (12, '4. Harta echipei + Misiunea imposibilă', 'Post-it-uri pe 6 zone de puncte forte; studiu de caz: de ce resurse are nevoie echipa.', 'Anexele 4, 5'),
        (8, '5. Mingea complimentelor', 'Feedback pozitiv concret, dat și primit.', 'Minge, Anexa 6'),
        (8, '6. Cardul de agent + biletul de ieșire', 'Integrare și reflecție; cardurile rămân la facilitator pentru Ședința 2.', 'Anexele 7, 9'),
    ]),
    ('ȘEDINȚA 1 – VARIANTA SCURTĂ (40 min) – Grupa 2, 09.10.2026, 10:15–10:55',
     'Aceleași obiective; se omite Dosarul secret și Misiunea imposibilă, iar energizarea se scurtează.', [
        (3, '0. Recrutarea agenților', 'Ecusoane și reguli, foarte pe scurt.', 'Anexa 1'),
        (4, '1. Schimbă locul dacă...', 'Doar rundele 2 și 3 (ce știu să fac, curaj și viitor).', '–'),
        (10, '2. Detectivii de puncte forte', 'Interviu în perechi, câte 4 minute pe rol; 2 rapoarte voluntare.', 'Anexa 3 (Fișa 1)'),
        (8, '3. Harta echipei', 'Post-it-uri pe cele 6 zone și o discuție scurtă despre diversitate.', 'Anexa 4'),
        (8, '4. Mingea complimentelor', 'Fiecare primește un compliment; regula „spun doar Mulțumesc”.', 'Minge, Anexa 6'),
        (7, '5. Cardul de agent + biletul de ieșire', 'Card completat pe scurt; biletul de ieșire predat la plecare.', 'Anexele 7, 9'),
    ]),
    ('ȘEDINȚA 2 – „Ce mă motivează și ce vreau să dezvolt?” (60 min) – ambele grupe, 29.10.2026 | tema creativă: LEVEL UP',
     'Scop: explorarea surselor de motivație și legarea punctelor forte de un obiectiv mic, realist, de 30 de zile.', [
        (4, '0. Reconectare – bateria mea', 'Cardurile de agent sunt returnate; check-in 1–5.', 'Carduri de agent'),
        (7, '1. Colțurile motivației', 'Elevii aleg, prin mișcare, motivația principală și pe cea de rezervă.', 'Anexa 1'),
        (7, '2. Motivația se schimbă', 'Situații pe grupe; combustibil și frâne; Fișa 3.', 'Anexele 2, 3'),
        (9, '3. Superputerea mea', 'Un punct forte transformat în „superputere”; kryptonita; trailerul eroului.', 'Anexa 4 (Fișa 4)'),
        (13, '4. Provocarea LEVEL UP 30', 'Transformarea unei dorințe în obiectiv mic, clar și posibil; semnul de carte pentru 30 de zile.', 'Anexele 5, 6'),
        (8, '5. Licitația super-echipei', '10 monede împărțite între 8 calități; licitație live și discuție.', 'Anexa 7'),
        (7, '6. Scrisoare către mine, peste un an', 'Scrisoare sigilată, păstrată de facilitator și returnată ulterior.', 'Anexa 8, plicuri'),
        (5, '7. Încheiere – un cuvânt + diplome', 'Fișa de feedback, runda „un cuvânt”, diplome „Agent LEVEL UP”.', 'Anexele 9, 10'),
    ]),
]
r = 4
for head, scop, rows in blocks:
    wt.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    c = wt.cell(row=r, column=1, value=head)
    c.font = f(True, 'FFFFFF', 11)
    c.fill = fill(PURPLE)
    c.alignment = Alignment(wrap_text=True, vertical='center')
    wt.row_dimensions[r].height = 22
    r += 1
    wt.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    c = wt.cell(row=r, column=1, value=scop)
    c.font = f(italic=True)
    c.fill = fill(LAV)
    r += 1
    hdr(wt, r, ['Minutul de început', 'Durata (min)', 'Activitatea', 'Ce fac elevii / scopul activității', 'Materiale (anexe)'])
    r += 1
    first = r
    for k, (dur, act, desc, mat) in enumerate(rows):
        cell(wt, r, 1, 0 if k == 0 else '=A%d+B%d' % (r - 1, r - 1), align='center')
        cell(wt, r, 2, dur, align='center', bg=YELLOW)
        cell(wt, r, 3, act, True)
        cell(wt, r, 4, desc)
        cell(wt, r, 5, mat)
        r += 1
    cell(wt, r, 1, 'TOTAL', True, bg=GREY, align='center')
    cell(wt, r, 2, '=SUM(B%d:B%d)' % (first, r - 1), True, bg=GREY, align='center')
    for col in (3, 4, 5):
        cell(wt, r, col, '', bg=GREY)
    r += 3
page(wt)

# ================================================================== 4. Elevi
we = wb.create_sheet('Elevi pe grupe')
title(we, 'Elevii implicați în proiectul UNIC 2 – Liceul Tehnologic „Vlădeasa” Huedin',
      'Date cu caracter personal ale unor minori – se utilizează doar în scopul proiectului. Prezență: P = prezent, A = absent.')
hdr(we, 4, ['Nr.', 'Numele și prenumele', 'Clasa', 'Grupa', 'Prezență Ședința 1', 'Prezență Ședința 2', 'Observații'],
    [5, 36, 9, 10, 12, 12, 34])
dvp = DataValidation(type='list', formula1='"P,A"', allow_blank=True)
we.add_data_validation(dvp)
for i, (nume, clasa) in enumerate(elevi):
    r = 5 + i
    cell(we, r, 1, i + 1, align='center')
    cell(we, r, 2, nume)
    cell(we, r, 3, clasa, align='center')
    cell(we, r, 4, "=IFERROR(INDEX('Sinteză'!$B$%d:$B$%d,MATCH(C%d,'Sinteză'!$A$%d:$A$%d,0)),\"\")"
         % (CLS_FIRST, CLS_LAST, r, CLS_FIRST, CLS_LAST), align='center')
    cell(we, r, 5, None, bg=YELLOW, align='center')
    cell(we, r, 6, None, bg=YELLOW, align='center')
    cell(we, r, 7, None)
    dvp.add('E%d:F%d' % (r, r))
we.freeze_panes = 'C5'
we.auto_filter.ref = 'A4:G%d' % (4 + len(elevi))
page(we, landscape=False)

# ================================================================== 5. Programare FC2 (Anexa 12)
wa = wb.create_sheet('Programare FC2 (Anexa 12)')
title(wa, 'Programarea Facilitatorului comunitar 2 în Anexa 12 (V2) – octombrie 2026, SA5.3',
      'Extras din Anexa 12; toate sesiunile FC2 din octombrie, inclusiv cele de la Turda.')
hdr(wa, 4, ['Nr.', 'Data', 'Ora început', 'Ora sfârșit', 'Durata (min)', 'Locația', 'GT declarat', 'Huedin?'], [5, 12, 10, 10, 10, 60, 11, 10])
fc2 = [
    (date(2026, 10, 2), time(11, 0), time(12, 0), 'Colegiul Tehnic Turda, Structura Poiana, str. Câmpiei nr. 51, Turda', 3),
    (date(2026, 10, 9), time(10, 15), time(10, 55), 'Liceul Tehnologic „Vlădeasa”, Huedin', 2),
    (date(2026, 10, 9), time(11, 55), time(12, 55), 'Liceul Tehnologic „Vlădeasa”, Huedin', 3),
    (date(2026, 10, 12), time(11, 0), time(12, 0), 'Colegiul Tehnic Turda, Structura Poiana, str. Câmpiei nr. 51, Turda', 3),
    (date(2026, 10, 29), time(11, 0), time(12, 0), 'Liceul Tehnologic „Vlădeasa”, Huedin', 3),
    (date(2026, 10, 29), time(12, 0), time(13, 0), 'Liceul Tehnologic „Vlădeasa”, Huedin', 3),
]
for i, (dt, t1, t2, loc, gt) in enumerate(fc2):
    r = 5 + i
    cell(wa, r, 1, i + 1, align='center')
    cell(wa, r, 2, dt, num='dd.mm.yyyy', align='center')
    cell(wa, r, 3, t1, num='hh:mm', align='center')
    cell(wa, r, 4, t2, num='hh:mm', align='center')
    cell(wa, r, 5, '=ROUND((D%d-C%d)*1440,0)' % (r, r), align='center')
    cell(wa, r, 6, loc)
    cell(wa, r, 7, gt, align='center')
    cell(wa, r, 8, '=IF(ISNUMBER(SEARCH("Huedin",F%d)),"Da","Nu")' % r, align='center')
r = 5 + len(fc2)
for col in (1, 3, 4, 8):
    cell(wa, r, col, '', bg=GREY)
cell(wa, r, 2, 'TOTAL', True, bg=GREY)
cell(wa, r, 5, '=SUM(E5:E%d)' % (r - 1), True, bg=GREY, align='center')
cell(wa, r, 6, '=COUNTIF(H5:H%d,"Da")&" sesiuni la Huedin, "&COUNTIF(H5:H%d,"Nu")&" la Turda; total ore: "&TEXT(E%d/60,"0.00")'
     % (r - 1, r - 1, r), True, bg=GREY)
cell(wa, r, 7, '=SUM(G5:G%d)' % (r - 1), True, bg=GREY, align='center')
page(wa)

# ================================================================== 6. Observații
wo = wb.create_sheet('Observații')
title(wo, 'Analiza programării FC2 din Anexa 12 (V2) – constatări și propuneri')
hdr(wo, 4, ['Nr.', 'Constatare', 'Detalii', 'Propunere'], [5, 34, 62, 62])
obs = [
    ('Programul FC2 în octombrie',
     '6 sesiuni SA5.3: 2 la Turda – Structura Poiana (02.10, 12.10) și 4 la Huedin (09.10 și 29.10, câte 2 sesiuni pe zi). '
     'La Huedin: 3 h 40 min cu grupul țintă.',
     'La Huedin, cele 4 sesiuni permit atelierul complet de 2 ședințe pentru 2 grupe: Ședința 1 pe 09.10, Ședința 2 pe 29.10.'),
    ('Intervalul 09.10, 10:15–10:55 are doar 40 de minute',
     'Scenariul ședințelor este proiectat pentru 60 de minute.',
     'Se aplică varianta scurtă de 40 de minute a Ședinței 1 (foaia „Tematici și activități”), cu grupa mai mică (Grupa 2). '
     'Alternativ: extinderea intervalului la 60 de minute și transmiterea unei versiuni noi a Anexei 12.'),
    ('GT declarat mult mai mic decât numărul de elevi repartizați',
     'Anexa 12 declară 2–3 elevi pe sesiune (11 în total la Huedin), iar repartizarea școlii are 24 de elevi: '
     '15 în clasele a VII-a și 9 în clasele a VIII-a.',
     'Dacă participă toți elevii din grupă, actualizați Anexa 12 (versiunea 3) cu numărul real, cu cel puțin o zi înainte de activitate, '
     'conform declarației din anexă. Lista de prezență trebuie să corespundă cu GT raportat.'),
    ('Aceleași intervale la doi experți, 29.10',
     'Intervalele 29.10, 11:00–12:00 și 12:00–13:00 la Liceul „Vlădeasa” apar în Anexa 12 atât la Facilitatorul comunitar 2, '
     'cât și la Expertul comunicare (Moraru Georgia).',
     'Clarificați cu managerul de proiect dacă este o activitate comună (de exemplu, Expertul comunicare documentează atelierul) sau o '
     'dublare. Evitați raportarea acelorași ore și a aceluiași GT de două ori.'),
    ('Numărul și data documentului',
     'Anexa 12 poartă „Nr. 2026092503 / 25.10.2026”; data este ulterioară majorității activităților planificate pentru octombrie.',
     'Verificați data: numărul de înregistrare sugerează 25.09.2026.'),
    ('Clasa a VIII-A are un singur elev',
     'Un elev nu poate forma singur o grupă de lucru.',
     'Elevul este inclus în Grupa 2, împreună cu clasa a VIII-B (9 elevi în total).'),
    ('Durata ședinței 2 pentru Grupa 2',
     'Grupa 2 face Ședința 1 în varianta scurtă, deci nu parcurge Dosarul secret și Misiunea imposibilă.',
     'Pe 29.10, la Reconectare (primele 4 minute), se poate face un mini-„Dosar secret” cu 3–4 bilețele.'),
]
for i, (a, b, c) in enumerate(obs):
    r = 5 + i
    cell(wo, r, 1, i + 1, align='center')
    cell(wo, r, 2, a, True)
    cell(wo, r, 3, b)
    cell(wo, r, 4, c)
    wo.row_dimensions[r].height = 64
page(wo)

wb.save(OUT)
print('OK', OUT, len(elevi), 'elevi', CLASE)
