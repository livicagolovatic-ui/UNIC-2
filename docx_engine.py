# -*- coding: utf-8 -*-
import docx, re
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = docx.Document()

# ---------- page setup ----------
s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.left_margin = s.right_margin = Cm(2.2)
s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0)

INK   = RGBColor(0x1A, 0x1A, 0x1A)
BLUE  = RGBColor(0x00, 0x3D, 0x82)
GREY  = RGBColor(0x5A, 0x5A, 0x5A)
RED   = RGBColor(0x9B, 0x1C, 0x1C)

st = doc.styles
n = st['Normal']
n.font.name = 'Calibri'; n.font.size = Pt(10.5); n.font.color.rgb = INK
n._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
n.paragraph_format.space_after = Pt(6)
n.paragraph_format.line_spacing = 1.12

def cfg(name, size, bold, color, before, after, keep=False):
    x = st[name]
    x.font.name = 'Calibri'; x.font.size = Pt(size); x.font.bold = bold
    x.font.color.rgb = color
    x.paragraph_format.space_before = Pt(before)
    x.paragraph_format.space_after  = Pt(after)
    x.paragraph_format.keep_with_next = keep
    return x

cfg('Heading 1', 17, True, BLUE, 20, 8, True)
cfg('Heading 2', 14, True, BLUE, 16, 6, True)
cfg('Heading 3', 12, True, INK,  12, 4, True)
cfg('Heading 4', 10.5, True, GREY, 10, 3, True)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)

def borders(tbl, color='BFBFBF', sz=4):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement('w:'+edge)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),str(sz))
        e.set(qn('w:space'),'0'); e.set(qn('w:color'),color)
        b.append(e)
    tblPr.append(b)

# ---------- inline markup: **bold**, //italic// ----------
TOK = re.compile(r'(\*\*.+?\*\*|//.+?//)', re.S)
def runs(p, text, size=None, color=None, italic=False, bold=False):
    for part in TOK.split(text):
        if not part: continue
        b, i = bold, italic
        if part.startswith('**') and part.endswith('**'): part, b = part[2:-2], True
        elif part.startswith('//') and part.endswith('//'): part, i = part[2:-2], True
        r = p.add_run(part); r.bold = b; r.italic = i
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color
    return p

CHARS = {'n': 0}
def clean(t):
    return re.sub(r'\s+', ' ', re.sub(r'\*\*|//', '', t)).strip()

# ---------- block helpers ----------
def H1(t): doc.add_paragraph(t, style='Heading 1')
def H2(t): doc.add_paragraph(t, style='Heading 2')
def H3(t): doc.add_paragraph(t, style='Heading 3')
def H4(t): doc.add_paragraph(t, style='Heading 4')

def P(t, count=True):
    p = doc.add_paragraph(); runs(p, t)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if count: CHARS['n'] += len(clean(t)) + 1
    return p

def SMALL(t):
    p = doc.add_paragraph(); runs(p, t, size=8.5, color=GREY, italic=True)
    p.paragraph_format.space_after = Pt(8)
    return p

def FLAG(t):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(6)
    runs(p, t, size=9, color=RED)
    return p

def BUL(items, count=True):
    for it in items:
        p = doc.add_paragraph(style='List Bullet'); runs(p, it)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.6)
        if count: CHARS['n'] += len(clean(it)) + 1

def NUM(items):
    for it in items:
        p = doc.add_paragraph(style='List Number'); runs(p, it)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.6)
        CHARS['n'] += len(clean(it)) + 1

LAST_Q = {'t': '(no question)'}

def Q(t):
    """A question of the online form."""
    LAST_Q['t'] = t
    tbl = doc.add_table(rows=1, cols=1); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0,0); shade(c, 'E8EEF7'); borders(tbl, 'C7D6EA', 4)
    p = c.paragraphs[0]; runs(p, t, size=10, color=BLUE, bold=True)
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    CHARS['n'] = 0          # start counting this answer

def COUNT(limit=5000, cut=None):
    n = CHARS['n']
    import os, sys
    if os.environ.get('FIELD_REPORT'):
        flag = 'OVER ' if n > limit else '     '
        print('%s%6d / %5d  %s' % (flag, n, limit, LAST_Q['t'][:95]), file=sys.stderr)
    p = doc.add_paragraph()
    over = n > limit
    mark = 'within limit' if not over else 'OVER by %s — see below' % f'{n - limit:,}'
    runs(p, '≈ %s characters (assumed form limit %s — %s)' % (f'{n:,}', f'{limit:,}', mark),
         size=8, color=GREY if not over else RED, italic=True)
    p.paragraph_format.space_after = Pt(2 if over else 10)
    if over and cut:
        q = doc.add_paragraph(); q.paragraph_format.left_indent = Cm(0.4)
        q.paragraph_format.space_after = Pt(10)
        runs(q, 'HOW TO CUT IT: ' + cut, size=8.5, color=RED)
    CHARS['n'] = 0

def TBL(rows, widths=None, header=True, small=False, count=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    t.autofit = True
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci)
            c.text = ''
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
            first = True
            for line in str(val).split('\n'):
                tgt = p if first else c.add_paragraph()
                if not first:
                    tgt.paragraph_format.space_before = Pt(0); tgt.paragraph_format.space_after = Pt(2)
                runs(tgt, line, size=8.5 if small else 9.5,
                     bold=(header and ri == 0), color=BLUE if (header and ri == 0) else None)
                first = False
            if header and ri == 0: shade(c, 'E8EEF7')
            if count: CHARS['n'] += len(clean(str(val))) + 1
    if widths:
        for ri in range(len(rows)):
            for ci, w in enumerate(widths):
                t.cell(ri, ci).width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def CALLOUT(title, body, fill='FFF6E5', line='E8C97A'):
    tbl = doc.add_table(rows=1, cols=1); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0,0); shade(c, fill); borders(tbl, line, 6)
    p = c.paragraphs[0]; runs(p, title, size=9.5, bold=True)
    p.paragraph_format.space_before = Pt(3)
    q = c.add_paragraph(); runs(q, body, size=9.5)
    q.paragraph_format.space_after = Pt(3)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def PAGEBREAK():
    doc.add_page_break()

